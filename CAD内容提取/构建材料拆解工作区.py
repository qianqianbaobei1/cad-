#!/usr/bin/env python3
"""
构建材料拆解工作区。

读取项目知识库，把 BOQ、图纸候选做法、材料提及、质量问题转换成可人工确认的拆解任务。
这一步不直接承诺“自动采购清单”，而是生成可追溯、可修正、可逐步补规则的数据底座。

用法：
    python3 构建材料拆解工作区.py
    python3 构建材料拆解工作区.py --project 宿州302
"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


SCRIPT_DIR = Path(__file__).parent
PROJECTS_DIR = SCRIPT_DIR / "projects"


def read_json(path: Path) -> Any:
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.write("\n")


def normalize_text(text: Any) -> str:
    return re.sub(r"\s+", "", str(text or "")).strip()


def build_material_master() -> list[dict[str, Any]]:
    return [
        {
            "material_id": "MAT-001",
            "standard_name": "腻子",
            "category": "涂装基层材料",
            "purchase_unit": "kg",
            "quantity_basis": "按涂刷面积折算，需确认遍数、理论用量、基层条件",
            "required_specs": ["内/外墙类型", "耐水/普通", "遍数", "理论用量kg/m²/遍"],
            "aliases": ["内墙腻子", "耐水腻子", "腻子粉"],
            "status": "seed",
        },
        {
            "material_id": "MAT-002",
            "standard_name": "乳胶漆",
            "category": "面层涂料",
            "purchase_unit": "kg",
            "quantity_basis": "按涂刷面积折算，需确认底漆/面漆、遍数、理论涂布率",
            "required_specs": ["底漆/面漆", "内/外墙", "遍数", "颜色", "理论涂布率"],
            "aliases": ["内墙乳胶漆", "涂料", "面漆", "底漆"],
            "status": "seed",
        },
        {
            "material_id": "MAT-003",
            "standard_name": "石膏找平材料",
            "category": "找平材料",
            "purchase_unit": "kg",
            "quantity_basis": "按面积和厚度折算，需确认厚度、容重、损耗",
            "required_specs": ["厚度mm", "材料类型", "容重kg/m³"],
            "aliases": ["石膏", "粉刷石膏", "石膏找平"],
            "status": "seed",
        },
        {
            "material_id": "MAT-004",
            "standard_name": "防滑地砖",
            "category": "块料面层",
            "purchase_unit": "片/m²",
            "quantity_basis": "优先按 BOQ 面积，结合规格、铺贴方式和损耗折算采购量",
            "required_specs": ["长宽规格", "厚度", "颜色/型号", "防滑等级", "损耗率"],
            "aliases": ["地砖", "地面砖", "面砖", "防滑面砖"],
            "status": "seed",
        },
        {
            "material_id": "MAT-005",
            "standard_name": "墙面砖",
            "category": "块料面层",
            "purchase_unit": "片/m²",
            "quantity_basis": "优先按 BOQ 面积，结合规格、排版和损耗折算采购量",
            "required_specs": ["长宽规格", "厚度", "颜色/型号", "阳角收口方式", "损耗率"],
            "aliases": ["墙砖", "墙面砖", "面砖"],
            "status": "seed",
        },
        {
            "material_id": "MAT-006",
            "standard_name": "水泥砂浆",
            "category": "粘结/找平材料",
            "purchase_unit": "m³",
            "quantity_basis": "按面积、厚度、配合比折算，需确认砂浆类型和厚度",
            "required_specs": ["配合比", "厚度mm", "用途", "损耗率"],
            "aliases": ["干硬性水泥砂浆", "砂浆", "1:3水泥砂浆", "1:2水泥砂浆"],
            "status": "seed",
        },
        {
            "material_id": "MAT-007",
            "standard_name": "美缝剂",
            "category": "嵌缝材料",
            "purchase_unit": "支/kg",
            "quantity_basis": "按砖规格、缝宽、缝深和面积折算",
            "required_specs": ["缝宽mm", "缝深mm", "颜色", "砖规格"],
            "aliases": ["勾缝剂", "填缝剂", "擦缝", "美缝剂"],
            "status": "seed",
        },
        {
            "material_id": "MAT-008",
            "standard_name": "纸面石膏板",
            "category": "吊顶板材",
            "purchase_unit": "张/m²",
            "quantity_basis": "按吊顶展开面积、板材规格、层数、损耗折算",
            "required_specs": ["厚度mm", "板材规格", "层数", "防水/普通"],
            "aliases": ["石膏板", "纸面石膏板", "防水石膏板"],
            "status": "seed",
        },
        {
            "material_id": "MAT-009",
            "standard_name": "轻钢龙骨",
            "category": "吊顶龙骨",
            "purchase_unit": "m",
            "quantity_basis": "按吊顶面积和龙骨间距折算，需确认主/副龙骨规格和间距",
            "required_specs": ["主龙骨规格", "副龙骨规格", "龙骨间距", "吊杆间距"],
            "aliases": ["龙骨", "轻钢龙骨", "主龙骨", "副龙骨"],
            "status": "seed",
        },
        {
            "material_id": "MAT-010",
            "standard_name": "成品栏杆",
            "category": "成品构件",
            "purchase_unit": "m",
            "quantity_basis": "按 BOQ 长度，结合高度、材质、立杆间距和安装方式确认",
            "required_specs": ["高度", "材质", "表面处理", "立杆间距", "扶手规格"],
            "aliases": ["栏杆", "扶手栏杆", "护窗栏杆", "阳台栏杆", "屋面栏杆"],
            "status": "seed",
        },
        {
            "material_id": "MAT-011",
            "standard_name": "成品门窗",
            "category": "成品构件",
            "purchase_unit": "m²/樘",
            "quantity_basis": "按 BOQ 面积或樘数，结合门窗表确认洞口尺寸、系列、五金、玻璃",
            "required_specs": ["系列", "开启方式", "尺寸", "玻璃配置", "五金", "防火等级"],
            "aliases": ["窗", "门", "塑钢平开窗", "彩塑推拉门", "钢制乙级防火门", "百叶窗"],
            "status": "seed",
        },
    ]


def build_aliases(materials: list[dict[str, Any]]) -> list[dict[str, Any]]:
    aliases = []
    for material in materials:
        names = [material["standard_name"], *material.get("aliases", [])]
        for alias in dict.fromkeys(names):
            aliases.append({
                "alias": alias,
                "material_id": material["material_id"],
                "standard_name": material["standard_name"],
                "match_type": "exact_or_keyword_seed",
                "status": "seed",
            })
    return sorted(aliases, key=lambda x: (x["material_id"], x["alias"]))


def build_decomposition_rules() -> list[dict[str, Any]]:
    return [
        {
            "rule_id": "RULE-PAINT-CEILING",
            "name": "天花油漆拆解",
            "match_keywords": ["天花油漆", "吊顶油漆", "顶棚乳胶漆"],
            "output_material_ids": ["MAT-001", "MAT-002"],
            "optional_material_ids": ["MAT-003"],
            "quantity_level": "B",
            "required_inputs": ["涂刷面积", "腻子遍数", "乳胶漆底/面漆遍数", "理论用量"],
            "risk_notes": ["BOQ 只有面积，不能直接得出 kg；必须有材料品牌或理论用量参数"],
        },
        {
            "rule_id": "RULE-PAINT-WALL-CEILING",
            "name": "墙面天花油漆拆解",
            "match_keywords": ["墙面天花油漆", "墙面、天花油漆", "墙面及天花油漆"],
            "output_material_ids": ["MAT-001", "MAT-002"],
            "optional_material_ids": ["MAT-003"],
            "quantity_level": "B",
            "required_inputs": ["墙面/天花面积拆分", "腻子遍数", "乳胶漆遍数", "石膏找平厚度"],
            "risk_notes": ["一个 BOQ 同时覆盖墙面和天花时，需要拆分适用部位"],
        },
        {
            "rule_id": "RULE-TILE-FLOOR",
            "name": "地砖铺贴拆解",
            "match_keywords": ["地砖", "防滑地砖", "公共区域地砖"],
            "output_material_ids": ["MAT-004", "MAT-006", "MAT-007"],
            "optional_material_ids": [],
            "quantity_level": "B",
            "required_inputs": ["地砖规格", "砂浆厚度/粘结材料", "缝宽", "损耗率"],
            "risk_notes": ["BOQ 面积可作为主材基数，但规格、排版、损耗会影响采购数量"],
        },
        {
            "rule_id": "RULE-TILE-WALL",
            "name": "墙砖铺贴拆解",
            "match_keywords": ["墙砖", "墙面砖", "公共区域墙砖"],
            "output_material_ids": ["MAT-005", "MAT-006", "MAT-007"],
            "optional_material_ids": [],
            "quantity_level": "B",
            "required_inputs": ["墙砖规格", "粘结层做法", "阳角收口", "缝宽", "损耗率"],
            "risk_notes": ["墙砖规格可能有多种，需要按区域或图纸材料表确认"],
        },
        {
            "rule_id": "RULE-CEILING",
            "name": "吊顶系统拆解",
            "match_keywords": ["吊顶"],
            "output_material_ids": ["MAT-008", "MAT-009"],
            "optional_material_ids": [],
            "quantity_level": "C",
            "required_inputs": ["吊顶构造", "石膏板层数", "板材规格", "龙骨间距", "吊杆间距"],
            "risk_notes": ["只有吊顶面积时，无法可靠拆出龙骨、吊杆、螺丝等辅材"],
        },
        {
            "rule_id": "RULE-RAILING",
            "name": "栏杆成品拆解",
            "match_keywords": ["栏杆", "扶手", "护窗"],
            "output_material_ids": ["MAT-010"],
            "optional_material_ids": [],
            "quantity_level": "C",
            "required_inputs": ["栏杆高度", "材质", "节点详图", "表面处理", "安装方式"],
            "risk_notes": ["BOQ 长度可用，但成品深化参数不足时不能形成采购规格"],
        },
        {
            "rule_id": "RULE-DOOR-WINDOW",
            "name": "门窗成品拆解",
            "match_keywords": ["窗", "门", "百叶窗", "防火门"],
            "output_material_ids": ["MAT-011"],
            "optional_material_ids": [],
            "quantity_level": "C",
            "required_inputs": ["门窗表", "洞口尺寸", "系列", "开启方式", "玻璃/五金配置"],
            "risk_notes": ["面积或樘数不能替代门窗表，需按编号和尺寸归集"],
        },
    ]


def select_rule(boq_item: dict[str, Any], rules: list[dict[str, Any]]) -> dict[str, Any] | None:
    text = normalize_text(f"{boq_item.get('raw_name')} {boq_item.get('raw_description')}")
    if "墙面天花" in text or "墙面及天花" in text or "墙面、天花" in text:
        return next(r for r in rules if r["rule_id"] == "RULE-PAINT-WALL-CEILING")
    if "吊顶油漆" in text or "天花油漆" in text or "顶棚" in text and "乳胶漆" in text:
        return next(r for r in rules if r["rule_id"] == "RULE-PAINT-CEILING")
    if "地砖" in text:
        return next(r for r in rules if r["rule_id"] == "RULE-TILE-FLOOR")
    if "墙砖" in text:
        return next(r for r in rules if r["rule_id"] == "RULE-TILE-WALL")
    if "吊顶" in text:
        return next(r for r in rules if r["rule_id"] == "RULE-CEILING")
    if any(k in text for k in ["栏杆", "扶手", "护窗"]):
        return next(r for r in rules if r["rule_id"] == "RULE-RAILING")
    if any(k in text for k in ["窗", "门", "百叶"]):
        return next(r for r in rules if r["rule_id"] == "RULE-DOOR-WINDOW")
    return None


def extract_specs_from_text(text: str) -> list[str]:
    specs = []
    patterns = [
        r"\d+\s*[*xX×]\s*\d+",
        r"\d+(?:\.\d+)?\s*mm",
        r"\d+(?:\.\d+)?\s*毫米",
        r"\d+(?:\.\d+)?\s*厚",
        r"\d+\s*高度",
        r"\d+系列",
        r"[甲乙丙]级防火门",
        r"C\d+",
        r"\d+:\d+(?:\.\d+)?",
        r"[一二两三四五六七八九十\d]+遍",
    ]
    for pattern in patterns:
        for match in re.findall(pattern, text):
            specs.append(re.sub(r"\s+", "", match))
    return list(dict.fromkeys(specs))


def index_quality_issues(issues: list[dict[str, Any]]) -> tuple[dict[str, list[dict[str, Any]]], dict[str, list[dict[str, Any]]]]:
    by_boq = defaultdict(list)
    by_practice = defaultdict(list)
    for issue in issues:
        if issue.get("boq_item_id"):
            by_boq[issue["boq_item_id"]].append(issue)
        if issue.get("practice_id"):
            by_practice[issue["practice_id"]].append(issue)
    return by_boq, by_practice


def build_tasks(project_name: str) -> dict[str, Any]:
    project_root = PROJECTS_DIR / project_name
    kb_dir = project_root / "项目知识库"
    work_dir = project_root / "材料拆解工作区"

    boq_items = read_json(kb_dir / "boq_items.json")
    links = read_json(kb_dir / "boq_drawing_links.json")
    mentions = read_json(kb_dir / "material_mentions.json")
    issues = read_json(kb_dir / "quality_issues.json")

    materials = build_material_master()
    material_by_id = {m["material_id"]: m for m in materials}
    aliases = build_aliases(materials)
    rules = build_decomposition_rules()

    links_by_boq = defaultdict(list)
    for link in links:
        links_by_boq[link["boq_item_id"]].append(link)
    for item_links in links_by_boq.values():
        item_links.sort(key=lambda x: (-float(x.get("confidence", 0)), x.get("rank", 999)))

    mentions_by_practice = defaultdict(list)
    for mention in mentions:
        mentions_by_practice[mention["practice_id"]].append(mention)

    issues_by_boq, issues_by_practice = index_quality_issues(issues)

    tasks = []
    missing_inputs = []
    accepted_seed = []
    pending_seed = []

    for boq in boq_items:
        rule = select_rule(boq, rules)
        top_links = links_by_boq.get(boq["boq_item_id"], [])[:5]
        evidence_mentions = []
        evidence_issue_count = 0
        for link in top_links:
            evidence_mentions.extend(mentions_by_practice.get(link["practice_id"], [])[:6])
            evidence_issue_count += len(issues_by_practice.get(link["practice_id"], []))

        raw_text = f"{boq.get('raw_name', '')} {boq.get('raw_description', '')}"
        specs_found = extract_specs_from_text(raw_text)
        link_statuses = Counter(link.get("status") for link in top_links)
        has_high_link = bool(link_statuses.get("auto_candidate_high"))
        has_links = bool(top_links)

        if not rule:
            task_status = "blocking_no_rule"
            confidence = 0.2
            grade = "D"
            proposed_materials = []
            required_inputs = ["补充清单项拆解规则"]
        else:
            grade = rule["quantity_level"]
            confidence = 0.45
            if boq.get("quantity") and boq.get("unit"):
                confidence += 0.15
            if has_high_link:
                confidence += 0.2
            elif has_links:
                confidence += 0.08
            if specs_found:
                confidence += 0.08
            if evidence_issue_count:
                confidence -= 0.08
            confidence = round(max(0.15, min(confidence, 0.9)), 2)
            task_status = "ready_for_human_confirm" if grade in ["B", "C"] else "blocking_missing_data"
            required_inputs = list(rule["required_inputs"])

            proposed_materials = []
            for idx, material_id in enumerate(rule["output_material_ids"], start=1):
                material = material_by_id[material_id]
                proposed_materials.append({
                    "line_id": f"{boq['boq_item_id']}-M{idx:02d}",
                    "material_id": material_id,
                    "standard_name": material["standard_name"],
                    "role": "主材" if idx == 1 else "辅材/配套材料",
                    "purchase_unit": material["purchase_unit"],
                    "quantity_formula_status": "formula_template_only",
                    "quantity_formula": build_formula_template(boq, material, rule),
                    "loss_rate_status": "needs_rule_or_manual_input",
                    "specs_found_in_boq": specs_found,
                    "required_specs": material["required_specs"],
                    "evidence_material_mentions": compact_mentions(evidence_mentions, material["standard_name"]),
                    "status": "pending_confirm",
                })

        missing = build_missing_inputs(boq, rule, top_links, specs_found, evidence_issue_count)
        for item in missing:
            missing_inputs.append(item)

        task = {
            "task_id": f"TASK-{boq['boq_item_id'].split('-')[-1]}",
            "boq_item_id": boq["boq_item_id"],
            "boq_name": boq.get("raw_name"),
            "boq_description": boq.get("raw_description"),
            "boq_unit": boq.get("unit"),
            "boq_quantity": boq.get("quantity"),
            "boq_category": boq.get("category"),
            "matched_rule_id": rule["rule_id"] if rule else None,
            "matched_rule_name": rule["name"] if rule else None,
            "quantity_grade": grade,
            "confidence": confidence,
            "task_status": task_status,
            "candidate_drawing_links": compact_links(top_links),
            "drawing_evidence_summary": {
                "candidate_count": len(links_by_boq.get(boq["boq_item_id"], [])),
                "top_link_statuses": dict(link_statuses),
                "top_practice_quality_issue_count": evidence_issue_count,
                "top_material_mentions": compact_mentions(evidence_mentions),
            },
            "proposed_materials": proposed_materials,
            "required_human_inputs": required_inputs,
            "missing_inputs": [x for x in missing if x["boq_item_id"] == boq["boq_item_id"]],
            "risk_notes": build_risk_notes(boq, rule, has_links, evidence_issue_count),
            "source_trace": {
                "boq_source": boq.get("source"),
                "knowledge_base": "项目知识库",
            },
        }
        tasks.append(task)

        if confidence >= 0.78 and grade == "B" and not missing:
            accepted_seed.append(task)
        else:
            pending_seed.append(task)

    output = {
        "work_dir": work_dir,
        "materials": materials,
        "aliases": aliases,
        "rules": rules,
        "tasks": tasks,
        "missing_inputs": missing_inputs,
        "accepted_seed": accepted_seed,
        "pending_seed": pending_seed,
    }
    return output


def build_formula_template(boq: dict[str, Any], material: dict[str, Any], rule: dict[str, Any]) -> str:
    name = material["standard_name"]
    unit = boq.get("unit") or "BOQ单位"
    qty = boq.get("quantity")
    qty_text = str(qty) if qty not in [None, ""] else "BOQ工程量"

    if name in ["防滑地砖", "墙面砖", "纸面石膏板", "成品栏杆", "成品门窗"]:
        return f"采购量 = {qty_text}{unit} × 损耗率/排版系数；需按规格换算采购单位"
    if name in ["腻子", "乳胶漆"]:
        return f"采购量 = {qty_text}{unit} × 遍数 × 理论用量 × 损耗率"
    if name in ["水泥砂浆", "石膏找平材料"]:
        return f"采购量 = {qty_text}{unit} × 厚度m × 材料用量系数 × 损耗率"
    if name == "美缝剂":
        return f"采购量 = {qty_text}{unit} × 单位面积缝长 × 缝宽 × 缝深 × 损耗率"
    if name == "轻钢龙骨":
        return f"采购量 = {qty_text}{unit} × 龙骨排布系数；需确认主副龙骨间距"
    return f"采购量 = {qty_text}{unit} × 规则系数 × 损耗率"


def compact_links(links: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {
            "link_id": link.get("link_id"),
            "practice_id": link.get("practice_id"),
            "building": link.get("building"),
            "practice_category": link.get("practice_category"),
            "practice_code": link.get("practice_code"),
            "location": link.get("location"),
            "confidence": link.get("confidence"),
            "status": link.get("status"),
            "match_basis": link.get("match_basis", []),
        }
        for link in links
    ]


def compact_mentions(mentions: list[dict[str, Any]], material_name: str | None = None) -> list[dict[str, Any]]:
    results = []
    seen = set()
    for mention in mentions:
        if material_name and material_name not in mention.get("standard_material_name", ""):
            continue
        key = (mention.get("standard_material_name"), mention.get("raw_text"))
        if key in seen:
            continue
        seen.add(key)
        results.append({
            "material_name": mention.get("standard_material_name"),
            "raw_text": mention.get("raw_text"),
            "spec": mention.get("spec"),
            "quality": mention.get("quality"),
            "source": mention.get("source"),
        })
        if len(results) >= 8:
            break
    return results


def build_missing_inputs(
    boq: dict[str, Any],
    rule: dict[str, Any] | None,
    top_links: list[dict[str, Any]],
    specs_found: list[str],
    evidence_issue_count: int,
) -> list[dict[str, Any]]:
    missing = []
    boq_id = boq["boq_item_id"]

    if not rule:
        missing.append({
            "missing_id": f"MISS-{boq_id}-RULE",
            "boq_item_id": boq_id,
            "type": "decomposition_rule",
            "severity": "blocking",
            "description": "缺少该清单项对应的材料拆解规则",
            "suggested_owner": "产品/造价专家",
        })
        return missing

    text = normalize_text(f"{boq.get('raw_name', '')} {boq.get('raw_description', '')}")
    for required in rule["required_inputs"]:
        if "面积" in required and boq.get("quantity") and boq.get("unit") == "m²":
            continue
        if "长度" in required and boq.get("quantity") and boq.get("unit") == "m":
            continue
        if "腻子遍数" in required and re.search(r"腻子[一二两三四五六七八九十\d]+遍|[一二两三四五六七八九十\d]+遍腻子", text):
            continue
        if "乳胶漆" in required and re.search(r"乳胶漆.*[一二两三四五六七八九十\d]+遍|[一二两三四五六七八九十\d]+遍.*乳胶漆", text):
            continue
        if "石膏找平厚度" in required and re.search(r"\d+(?:\.\d+)?(?:mm|毫米).*石膏|石膏.*\d+(?:\.\d+)?(?:mm|毫米)", text):
            continue
        if any(token in required for token in ["规格", "高度", "系列", "防火等级"]) and specs_found:
            continue
        missing.append({
            "missing_id": f"MISS-{boq_id}-{len(missing)+1:02d}",
            "boq_item_id": boq_id,
            "type": "calculation_or_spec_input",
            "severity": "high" if rule["quantity_level"] == "C" else "medium",
            "description": required,
            "suggested_owner": "人工确认/项目资料补充",
        })

    if not top_links and boq.get("category") not in ["栏杆材料", "门窗材料"]:
        missing.append({
            "missing_id": f"MISS-{boq_id}-DRAWING",
            "boq_item_id": boq_id,
            "type": "drawing_evidence",
            "severity": "high",
            "description": "没有匹配到图纸做法候选，不能用图纸补充规格和做法",
            "suggested_owner": "技术解析/人工关联",
        })

    if evidence_issue_count:
        missing.append({
            "missing_id": f"MISS-{boq_id}-QUALITY",
            "boq_item_id": boq_id,
            "type": "source_quality",
            "severity": "medium",
            "description": f"候选图纸做法存在 {evidence_issue_count} 个质量问题，需要复核后再参与计算",
            "suggested_owner": "人工复核/图纸解析优化",
        })

    return missing


def build_risk_notes(boq: dict[str, Any], rule: dict[str, Any] | None, has_links: bool, evidence_issue_count: int) -> list[str]:
    notes = []
    if rule:
        notes.extend(rule.get("risk_notes", []))
    else:
        notes.append("当前系统无法识别该清单项的拆解逻辑")
    if not has_links and boq.get("category") not in ["栏杆材料", "门窗材料"]:
        notes.append("未建立清单与图纸做法的对应关系，图纸暂时不能发挥补充作用")
    if evidence_issue_count:
        notes.append("图纸候选做法存在 OCR/拼接/占位问题，不能作为 A 级计算依据")
    if boq.get("unit") in ["元/m²", "元/㎡"] or str(boq.get("raw_unit", "")).startswith("元/"):
        notes.append("原清单单位包含价格口径，已尝试归一为工程量单位，需确认是否符合业务口径")
    return list(dict.fromkeys(notes))


def write_review_workbook(path: Path, tasks: list[dict[str, Any]], missing_inputs: list[dict[str, Any]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "拆解任务队列"
    headers = [
        "任务ID", "BOQ ID", "清单名称", "工程量", "单位", "匹配规则", "等级", "置信度",
        "状态", "候选图纸数", "建议材料", "需人工确认", "主要风险",
    ]
    ws.append(headers)
    for task in tasks:
        ws.append([
            task["task_id"],
            task["boq_item_id"],
            task["boq_name"],
            task["boq_quantity"],
            task["boq_unit"],
            task["matched_rule_name"],
            task["quantity_grade"],
            task["confidence"],
            task["task_status"],
            task["drawing_evidence_summary"]["candidate_count"],
            "、".join(m["standard_name"] for m in task["proposed_materials"]),
            "；".join(task["required_human_inputs"]),
            "；".join(task["risk_notes"]),
        ])

    ws2 = wb.create_sheet("缺失资料清单")
    headers2 = ["缺失ID", "BOQ ID", "类型", "严重度", "说明", "建议负责人"]
    ws2.append(headers2)
    for item in missing_inputs:
        ws2.append([
            item["missing_id"], item["boq_item_id"], item["type"],
            item["severity"], item["description"], item["suggested_owner"],
        ])

    ws3 = wb.create_sheet("候选材料明细")
    headers3 = ["任务ID", "BOQ ID", "清单名称", "材料ID", "标准材料名", "角色", "采购单位", "公式模板", "状态"]
    ws3.append(headers3)
    for task in tasks:
        for material in task["proposed_materials"]:
            ws3.append([
                task["task_id"], task["boq_item_id"], task["boq_name"],
                material["material_id"], material["standard_name"], material["role"],
                material["purchase_unit"], material["quantity_formula"], material["status"],
            ])

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        for column_cells in sheet.columns:
            length = max(len(str(cell.value or "")) for cell in column_cells[:80])
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(length + 2, 10), 50)
        sheet.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", default="宿州302")
    args = parser.parse_args()

    output = build_tasks(args.project)
    work_dir: Path = output["work_dir"]

    write_json(work_dir / "materials_master.json", output["materials"])
    write_json(work_dir / "material_aliases.json", output["aliases"])
    write_json(work_dir / "decomposition_rules.json", output["rules"])
    write_json(work_dir / "拆解任务队列.json", output["tasks"])
    write_json(work_dir / "缺失资料清单.json", output["missing_inputs"])
    write_json(work_dir / "可采纳清单_种子.json", output["accepted_seed"])
    write_json(work_dir / "待人工确认清单.json", output["pending_seed"])

    write_review_workbook(work_dir / "材料拆解工作区_人工确认.xlsx", output["tasks"], output["missing_inputs"])

    manifest = {
        "project": args.project,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "work_dir": "材料拆解工作区",
        "files": {
            "materials_master": "材料拆解工作区/materials_master.json",
            "material_aliases": "材料拆解工作区/material_aliases.json",
            "decomposition_rules": "材料拆解工作区/decomposition_rules.json",
            "task_queue": "材料拆解工作区/拆解任务队列.json",
            "missing_inputs": "材料拆解工作区/缺失资料清单.json",
            "accepted_seed": "材料拆解工作区/可采纳清单_种子.json",
            "pending_seed": "材料拆解工作区/待人工确认清单.json",
            "review_workbook": "材料拆解工作区/材料拆解工作区_人工确认.xlsx",
        },
        "counts": {
            "materials_master": len(output["materials"]),
            "material_aliases": len(output["aliases"]),
            "decomposition_rules": len(output["rules"]),
            "tasks": len(output["tasks"]),
            "missing_inputs": len(output["missing_inputs"]),
            "accepted_seed": len(output["accepted_seed"]),
            "pending_seed": len(output["pending_seed"]),
        },
        "important_note": "当前输出是拆解工作区，不是最终采购清单。所有 formula_template_only 和 pending_confirm 项必须经过规则完善或人工确认。",
    }
    write_json(work_dir / "manifest.json", manifest)

    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
