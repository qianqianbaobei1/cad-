#!/usr/bin/env python3
"""
构建项目级知识库。

把当前项目的 BOQ、CAD 做法表清洗结果整理成材料拆解可用的上下文：
1. boq_items.json            - 结构化 BOQ
2. practice_rules.json       - 项目做法规则
3. material_mentions.json    - 图纸中出现的材料提及
4. boq_drawing_links.json    - BOQ 与图纸做法的候选对应关系
5. quality_issues.json       - 资料/解析/规格质量问题

用法：
    python3 构建项目知识库.py
    python3 构建项目知识库.py --project 宿州302
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).parent
PROJECTS_DIR = SCRIPT_DIR / "projects"

SURFACE_BY_CATEGORY = {
    "地面做法": "地面",
    "楼面做法": "楼面",
    "内墙做法": "墙面",
    "外墙做法": "墙面",
    "顶棚做法": "天花",
}

BOQ_CATEGORY_RULES = [
    ("地面材料", ["地砖", "石材", "木地板", "地毯", "楼地面", "地面"]),
    ("墙面材料", ["墙砖", "墙面", "内墙", "外墙"]),
    ("天花材料", ["吊顶", "天花", "顶棚"]),
    ("栏杆材料", ["栏杆", "扶手", "护窗"]),
    ("门窗材料", ["门", "窗"]),
]

BOQ_CATEGORY_TO_PRACTICE = {
    "地面材料": ["地面做法", "楼面做法"],
    "墙面材料": ["内墙做法", "外墙做法"],
    "天花材料": ["顶棚做法"],
    "门窗材料": [],
    "栏杆材料": [],
    "复合材料": ["内墙做法", "外墙做法", "顶棚做法"],
    "未分类": [],
}

PRACTICE_CATEGORY_RULES = {
    "地面做法": ["地砖", "防滑地砖", "面砖", "石材", "木地板", "水泥砂浆", "防水涂料"],
    "楼面做法": ["地砖", "防滑地砖", "面砖", "石材", "木地板", "水泥砂浆", "防水涂料"],
    "内墙做法": ["墙砖", "面砖", "乳胶漆", "涂料", "腻子", "石膏", "砂浆"],
    "外墙做法": ["墙砖", "面砖", "涂料", "保温板", "锚栓", "镀锌钢丝网"],
    "顶棚做法": ["吊顶", "龙骨", "石膏板", "乳胶漆", "涂料", "腻子"],
}

MATERIAL_PATTERNS = [
    ("防滑地砖", r"防滑地砖|地砖|面砖"),
    ("墙面砖", r"墙砖|墙面砖|面砖"),
    ("JS聚合物水泥基防水涂料", r"JS.*防水涂料|聚合物水泥.*防水"),
    ("水泥砂浆", r"\d+厚1:\d+(?:\.\d+)?(?:水泥)?砂浆|干硬性水泥砂浆|水泥砂浆"),
    ("混合砂浆", r"混合砂浆"),
    ("细石混凝土", r"细石混凝土|细石砼"),
    ("混凝土", r"C\d+混凝土|C\d+砼"),
    ("水泥浆", r"水泥浆"),
    ("美缝剂", r"美缝剂|勾缝|擦缝"),
    ("腻子", r"腻子"),
    ("乳胶漆", r"乳胶漆"),
    ("涂料", r"涂料"),
    ("石膏", r"石膏"),
    ("石膏板", r"石膏板"),
    ("龙骨", r"龙骨"),
    ("保温板", r"保温板|聚苯乙烯保温隔声板"),
    ("锚栓", r"锚栓"),
    ("镀锌钢丝网", r"镀锌钢丝网|钢丝网"),
    ("界面剂", r"界面剂|界面处理剂|界面砂浆"),
    ("建筑胶", r"建筑胶"),
]

DIRTY_PATTERNS = [
    ("possible_ocr_typo", r"千硬性|中\d+@|%%|、、|1\.10mm|[,，]随打随抹平2、"),
    ("unknown_dimension", r"\bX厚|、、X厚|[，,]X厚"),
    ("merged_layers", r"防水涂料防水层.*水泥砂浆|水泥砂浆保护层.*JS"),
    ("placeholder_layer", r"预留装修面层|二次装修为准"),
]


@dataclass
class ProjectPaths:
    root: Path
    input_dir: Path
    extract_dir: Path
    kb_dir: Path
    output_dir: Path


def stable_id(*parts: Any, prefix: str = "") -> str:
    raw = "|".join(str(p) for p in parts)
    digest = hashlib.md5(raw.encode("utf-8")).hexdigest()[:10]
    return f"{prefix}{digest}" if prefix else digest


def read_json(path: Path) -> Any:
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.write("\n")


def project_paths(project_name: str) -> ProjectPaths:
    root = PROJECTS_DIR / project_name
    if not root.exists():
        raise FileNotFoundError(f"项目不存在: {root}")
    return ProjectPaths(
        root=root,
        input_dir=root / "原始输入",
        extract_dir=root / "提取结果" / "做法表",
        kb_dir=root / "项目知识库",
        output_dir=root / "输出产物",
    )


def normalize_unit(raw_unit: Any) -> str:
    text = str(raw_unit or "").strip()
    text = text.replace("㎡", "m²").replace("平方", "m²")
    if "m²" in text or "m2" in text or "平米" in text:
        return "m²"
    if "m" in text and "m²" not in text:
        return "m"
    if "个" in text or "只" in text:
        return "个"
    return text


def infer_boq_category(name: str, desc: str) -> tuple[str, list[str]]:
    haystack = f"{name} {desc}"
    if any(kw in haystack for kw in ["墙面天花", "墙面、天花", "墙面及天花"]):
        return "复合材料", BOQ_CATEGORY_TO_PRACTICE["复合材料"]
    if any(kw in haystack for kw in ["吊顶", "天花", "顶棚"]):
        return "天花材料", BOQ_CATEGORY_TO_PRACTICE["天花材料"]
    for category, keywords in BOQ_CATEGORY_RULES:
        if any(kw in haystack for kw in keywords):
            return category, BOQ_CATEGORY_TO_PRACTICE.get(category, [])
    return "未分类", []


def parse_references(text: str) -> list[dict[str, Any]]:
    refs = []
    for pattern, ref_type in [
        (r"详见[^，,；;。]*材料表", "material_table"),
        (r"详见[^，,；;。]*图纸", "drawing"),
        (r"详见[^，,；;。]*说明", "design_note"),
        (r"按[^，,；;。]*要求", "requirement"),
    ]:
        for m in re.finditer(pattern, text):
            refs.append({"type": ref_type, "text": m.group(0), "status": "unresolved"})
    return refs


def parse_boq(paths: ProjectPaths) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    boq_path = paths.input_dir / "BOQ_宿州302.xlsx"
    if not boq_path.exists():
        candidates = sorted(paths.input_dir.glob("*.xlsx"))
        if not candidates:
            return [], [{
                "issue_type": "missing_boq",
                "severity": "blocking",
                "message": "未找到 BOQ Excel",
                "source": str(paths.input_dir),
            }]
        boq_path = candidates[0]

    wb = load_workbook(boq_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))

    issues = []
    items = []

    for row_idx, row in enumerate(rows, start=1):
        seq = row[0] if len(row) > 0 else None
        if not isinstance(seq, int):
            continue
        name = str(row[1] or "").strip() if len(row) > 1 else ""
        desc = str(row[2] or "").strip() if len(row) > 2 else ""
        raw_unit = row[3] if len(row) > 3 else ""
        qty = row[15] if len(row) > 15 else None

        item_id = f"BOQ-{seq:03d}"
        unit = normalize_unit(raw_unit)
        category, target_practice_categories = infer_boq_category(name, desc)
        refs = parse_references(f"{name} {desc}")
        status = "parsed"

        if qty in (None, ""):
            status = "needs_review"
            issues.append({
                "issue_type": "boq_quantity_missing",
                "severity": "blocking",
                "boq_item_id": item_id,
                "message": "BOQ 缺工程量，不能自动计算采购量",
                "source": f"{boq_path.name}:row{row_idx}",
            })
        if not unit:
            status = "needs_review"
            issues.append({
                "issue_type": "boq_unit_missing",
                "severity": "high",
                "boq_item_id": item_id,
                "message": "BOQ 缺计量单位，需要人工确认",
                "source": f"{boq_path.name}:row{row_idx}",
            })

        items.append({
            "boq_item_id": item_id,
            "row_index": row_idx,
            "seq": seq,
            "raw_name": name,
            "raw_description": desc,
            "raw_unit": str(raw_unit or ""),
            "unit": unit,
            "quantity": qty,
            "category": category,
            "target_practice_categories": target_practice_categories,
            "standard_code": None,
            "code_match_status": "not_available",
            "references": refs,
            "status": status,
            "source": {
                "file": str(boq_path.relative_to(paths.root)),
                "sheet": ws.title,
                "row": row_idx,
            },
        })

    if not items:
        issues.append({
            "issue_type": "boq_no_items",
            "severity": "blocking",
            "message": "BOQ 中未识别到清单项",
            "source": str(boq_path),
        })

    return items, issues


def detect_quality(text: str, completeness: str = "") -> tuple[str, list[str]]:
    issue_codes = []
    for code, pattern in DIRTY_PATTERNS:
        if re.search(pattern, text):
            issue_codes.append(code)
    if completeness == "部分":
        issue_codes.append("partial_layer")
    if not text.strip():
        issue_codes.append("empty_layer")
    if issue_codes:
        return "needs_review", issue_codes
    return "valid", []


def extract_material_mentions(layer_text: str) -> list[dict[str, Any]]:
    mentions = []
    for standard_name, pattern in MATERIAL_PATTERNS:
        if re.search(pattern, layer_text):
            spec_parts = []
            for m in re.finditer(r"(\d+(?:\.\d+)?(?:~\d+(?:\.\d+)?)?厚|\d+×\d+|\d+\*\d+|C\d+|1:\d+(?:\.\d+)?|[AB]\d?级|II型|Ⅱ型)", layer_text):
                spec_parts.append(m.group(0))
            mentions.append({
                "raw_text": layer_text,
                "raw_material_name": standard_name,
                "standard_material_name": standard_name,
                "spec": "、".join(dict.fromkeys(spec_parts)),
                "match_method": "rule_keyword",
                "status": "standardized_by_seed_rule",
            })
    return mentions


def load_practice_results(paths: ProjectPaths) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    practice_rules = []
    material_mentions = []
    quality_issues = []

    for file_path in sorted(paths.extract_dir.glob("做法表_*.json")):
        data = read_json(file_path)
        building = data.get("building") or file_path.stem.replace("做法表_", "")
        relative_source = str(file_path.relative_to(paths.root))

        for category in SURFACE_BY_CATEGORY:
            items = data.get(category, [])
            if not isinstance(items, list):
                continue
            for item_index, item in enumerate(items):
                code = str(item.get("编号", "")).strip()
                location = str(item.get("部位", "")).strip()
                practice_id = f"PR-{stable_id(building, category, code, location, item_index)}"
                layers = []

                for layer_index, layer in enumerate(item.get("构造层次", []) or [], start=1):
                    if isinstance(layer, str):
                        layer_text = layer
                        completeness = ""
                        seq = layer_index
                    else:
                        layer_text = str(layer.get("做法", "")).strip()
                        completeness = str(layer.get("完整度", "")).strip()
                        seq = layer.get("序号") or layer_index

                    quality, issue_codes = detect_quality(layer_text, completeness)
                    layer_id = f"{practice_id}-L{layer_index:02d}"
                    layer_mentions = extract_material_mentions(layer_text)

                    for mention_index, mention in enumerate(layer_mentions, start=1):
                        mention_id = f"MM-{stable_id(layer_id, mention['standard_material_name'], mention_index)}"
                        material_mentions.append({
                            "mention_id": mention_id,
                            "practice_id": practice_id,
                            "layer_id": layer_id,
                            "building": building,
                            "category": category,
                            "surface": SURFACE_BY_CATEGORY[category],
                            "location": location,
                            **mention,
                            "quality": quality,
                            "source": {
                                "file": relative_source,
                                "practice_code": code,
                                "layer_seq": seq,
                            },
                        })

                    if issue_codes:
                        quality_issues.append({
                            "issue_id": f"QI-{stable_id(layer_id, ','.join(issue_codes))}",
                            "issue_type": "practice_layer_quality",
                            "severity": "medium" if quality == "needs_review" else "low",
                            "practice_id": practice_id,
                            "layer_id": layer_id,
                            "issue_codes": issue_codes,
                            "message": "做法层存在 OCR/拼接/完整度问题，需要人工复核",
                            "text": layer_text,
                            "source": {
                                "file": relative_source,
                                "practice_code": code,
                                "layer_seq": seq,
                            },
                        })

                    layers.append({
                        "layer_id": layer_id,
                        "seq": seq,
                        "text": layer_text,
                        "completeness": completeness or "未知",
                        "quality": quality,
                        "issue_codes": issue_codes,
                        "material_mention_ids": [
                            m["mention_id"]
                            for m in material_mentions
                            if m.get("layer_id") == layer_id
                        ],
                    })

                practice_rules.append({
                    "practice_id": practice_id,
                    "building": building,
                    "category": category,
                    "surface": SURFACE_BY_CATEGORY[category],
                    "code": code,
                    "location": location,
                    "layers": layers,
                    "source": {
                        "file": relative_source,
                        "item_index": item_index,
                    },
                    "status": "needs_review" if any(l["quality"] != "valid" for l in layers) else "parsed",
                })

    return practice_rules, material_mentions, quality_issues


def link_score(boq_item: dict[str, Any], practice: dict[str, Any], mentions: list[dict[str, Any]]) -> tuple[float, list[str]]:
    score = 0.0
    basis = []
    boq_text = f"{boq_item.get('raw_name', '')} {boq_item.get('raw_description', '')}"
    practice_text = " ".join(l.get("text", "") for l in practice.get("layers", []))
    location = practice.get("location", "")

    category = boq_item.get("category")
    target_categories = boq_item.get("target_practice_categories") or []
    if target_categories and practice.get("category") not in target_categories:
        return 0.0, []

    if category == "地面材料" and practice.get("category") in ["地面做法", "楼面做法"]:
        score += 0.30
        basis.append("BOQ类别为地面材料，做法类别为地面/楼面")
    elif category == "墙面材料" and practice.get("category") in ["内墙做法", "外墙做法"]:
        score += 0.30
        basis.append("BOQ类别为墙面材料，做法类别为内墙/外墙")
    elif category == "天花材料" and practice.get("category") == "顶棚做法":
        score += 0.30
        basis.append("BOQ类别为天花材料，做法类别为顶棚")
    elif category == "复合材料" and practice.get("category") in target_categories:
        score += 0.25
        basis.append("BOQ类别为墙面/天花复合项，做法类别在候选范围内")

    keyword_hits = []
    for kw in ["地砖", "防滑地砖", "墙砖", "面砖", "乳胶漆", "涂料", "腻子", "吊顶", "龙骨", "石膏板", "防水"]:
        if kw in boq_text and kw in practice_text:
            keyword_hits.append(kw)
    if keyword_hits:
        score += min(0.35, 0.12 * len(keyword_hits))
        basis.append(f"BOQ与做法层共同命中关键词：{','.join(keyword_hits)}")

    if "公共" in boq_text and any(kw in location for kw in ["公共", "电梯厅", "前室", "大堂", "楼梯间", "过道", "门厅"]):
        score += 0.20
        basis.append("BOQ为公共区域，做法部位属于公共/交通空间")
    if "卫生间" in boq_text and "卫生间" in location:
        score += 0.15
        basis.append("BOQ和做法部位均包含卫生间")
    if "厨房" in boq_text and "厨房" in location:
        score += 0.15
        basis.append("BOQ和做法部位均包含厨房")

    mention_names = {m["standard_material_name"] for m in mentions if m.get("practice_id") == practice.get("practice_id")}
    if mention_names:
        relevant = [m for m in mention_names if any(k in boq_text for k in [m, m.replace("防滑", "")])]
        if relevant:
            score += 0.15
            basis.append(f"材料提及与BOQ描述相关：{','.join(relevant)}")

    return min(score, 0.99), basis


def build_links(boq_items: list[dict[str, Any]], practice_rules: list[dict[str, Any]], material_mentions: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    links = []
    issues = []

    for boq in boq_items:
        candidates = []
        for practice in practice_rules:
            score, basis = link_score(boq, practice, material_mentions)
            if score >= 0.30:
                candidates.append((score, practice, basis))

        candidates.sort(key=lambda x: x[0], reverse=True)
        top = candidates[:20]
        if not top:
            issues.append({
                "issue_type": "boq_drawing_no_match",
                "severity": "high",
                "boq_item_id": boq["boq_item_id"],
                "message": "BOQ 未找到图纸做法候选，需要人工匹配",
            })
            continue

        for rank, (score, practice, basis) in enumerate(top, start=1):
            if score >= 0.75:
                status = "auto_candidate_high"
            elif score >= 0.55:
                status = "candidate_needs_review"
            else:
                status = "weak_candidate"

            links.append({
                "link_id": f"LINK-{stable_id(boq['boq_item_id'], practice['practice_id'])}",
                "boq_item_id": boq["boq_item_id"],
                "practice_id": practice["practice_id"],
                "building": practice["building"],
                "practice_category": practice["category"],
                "practice_code": practice["code"],
                "location": practice["location"],
                "confidence": round(score, 3),
                "rank": rank,
                "match_basis": basis,
                "status": status,
            })

    return links, issues


def build_summary(paths: ProjectPaths, outputs: dict[str, Any]) -> dict[str, Any]:
    return {
        "project": paths.root.name,
        "knowledge_base_dir": str(paths.kb_dir.relative_to(paths.root)),
        "files": {name: str((paths.kb_dir / filename).relative_to(paths.root)) for name, filename in outputs.items()},
        "next_use": [
            "材料拆解时先读取 boq_items.json 获取工程量和原始清单描述",
            "读取 boq_drawing_links.json 找到每条 BOQ 对应的候选图纸做法",
            "读取 practice_rules.json 和 material_mentions.json 作为拆料上下文",
            "读取 quality_issues.json，把脏数据和低置信项放入人工确认队列",
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", default="宿州302", help="项目名称，默认宿州302")
    args = parser.parse_args()

    paths = project_paths(args.project)
    paths.kb_dir.mkdir(parents=True, exist_ok=True)

    boq_items, boq_issues = parse_boq(paths)
    practice_rules, material_mentions, practice_issues = load_practice_results(paths)
    links, link_issues = build_links(boq_items, practice_rules, material_mentions)
    quality_issues = boq_issues + practice_issues + link_issues

    write_json(paths.kb_dir / "boq_items.json", boq_items)
    write_json(paths.kb_dir / "practice_rules.json", practice_rules)
    write_json(paths.kb_dir / "material_mentions.json", material_mentions)
    write_json(paths.kb_dir / "boq_drawing_links.json", links)
    write_json(paths.kb_dir / "quality_issues.json", quality_issues)

    manifest = build_summary(paths, {
        "boq_items": "boq_items.json",
        "practice_rules": "practice_rules.json",
        "material_mentions": "material_mentions.json",
        "boq_drawing_links": "boq_drawing_links.json",
        "quality_issues": "quality_issues.json",
    })
    manifest["counts"] = {
        "boq_items": len(boq_items),
        "practice_rules": len(practice_rules),
        "material_mentions": len(material_mentions),
        "boq_drawing_links": len(links),
        "quality_issues": len(quality_issues),
    }
    write_json(paths.kb_dir / "manifest.json", manifest)

    print("项目知识库已生成：")
    for key, val in manifest["counts"].items():
        print(f"  {key}: {val}")
    print(f"  输出目录: {paths.kb_dir}")


if __name__ == "__main__":
    main()
