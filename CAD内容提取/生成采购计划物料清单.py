#!/usr/bin/env python3
"""
生成最终交付的采购计划物料清单。

这个脚本只输出一份面向业务的物料清单 Excel，不输出过程表、任务表、问题表。

用法：
    python3 生成采购计划物料清单.py
    python3 生成采购计划物料清单.py --project 宿州302
"""

from __future__ import annotations

import argparse
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SCRIPT_DIR = Path(__file__).parent
PROJECTS_DIR = SCRIPT_DIR / "projects"


def read_json(path: Path) -> Any:
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def normalize_cn_number(text: str) -> str:
    mapping = {
        "一": "1", "二": "2", "两": "2", "三": "3", "四": "4", "五": "5",
        "六": "6", "七": "7", "八": "8", "九": "9", "十": "10",
    }
    for cn, num in mapping.items():
        text = text.replace(cn, num)
    return text


def unique_join(values: list[str]) -> str:
    cleaned = [str(v).strip() for v in values if str(v or "").strip()]
    return "、".join(dict.fromkeys(cleaned))


def extract_dimensions(text: str) -> str:
    matches = re.findall(r"\d+\s*[*xX×]\s*\d+", text)
    return unique_join([re.sub(r"\s+", "", x).replace("x", "*").replace("X", "*").replace("×", "*") for x in matches])


def extract_height(text: str) -> str:
    match = re.search(r"(\d+)\s*高度", text)
    return f"{match.group(1)}高度" if match else ""


def extract_paint_pass(text: str, material_name: str) -> str:
    text = normalize_cn_number(text)
    if material_name == "腻子":
        patterns = [
            r"腻子(\d+)遍",
            r"(\d+)遍腻子",
            r"腻子[^\d]{0,8}(\d+)遍",
        ]
    elif material_name == "乳胶漆":
        patterns = [
            r"乳胶漆(?:涂料)?(\d+)遍",
            r"乳胶漆",
        ]
    else:
        return ""

    found = []
    for pattern in patterns:
        for match in re.findall(pattern, text):
            if isinstance(match, tuple):
                match = next((m for m in match if m), "")
            if match:
                found.append(f"{match}遍" if match.isdigit() else "乳胶漆")
    if material_name == "乳胶漆" and any(item.endswith("遍") for item in found):
        found = [item for item in found if item.endswith("遍")]
    if material_name == "乳胶漆" and found == ["乳胶漆"]:
        return "乳胶漆，遍数待确认"
    return unique_join(found)


def extract_plaster_spec(text: str) -> str:
    match = re.search(r"(\d+(?:\.\d+)?)\s*(?:mm|毫米).*?石膏|石膏.*?(\d+(?:\.\d+)?)\s*(?:mm|毫米)", text)
    if not match:
        return ""
    value = next((x for x in match.groups() if x), "")
    return f"{value}mm石膏找平"


def extract_door_window_spec(text: str) -> str:
    text = re.sub(r"\s+", "", text)
    return text or "待确认"


def pick_spec(material: dict[str, Any], task: dict[str, Any]) -> str:
    name = material.get("standard_name", "")
    desc = f"{task.get('boq_name', '')} {task.get('boq_description', '')}"

    if name in ["腻子", "乳胶漆"]:
        return extract_paint_pass(desc, name) or "待确认"
    if name == "石膏找平材料":
        return extract_plaster_spec(desc) or "待确认"
    if name in ["防滑地砖", "墙面砖"]:
        return extract_dimensions(desc) or "待确认"
    if name in ["水泥砂浆"]:
        return "待确认"
    if name == "美缝剂":
        return "缝宽/颜色待确认"
    if name in ["纸面石膏板", "轻钢龙骨"]:
        return pick_spec_from_mentions(material, [name.replace("纸面", ""), "龙骨"]) or "待确认"
    if name == "成品栏杆":
        return extract_height(desc) or task.get("boq_description") or "待确认"
    if name == "成品门窗":
        return extract_door_window_spec(task.get("boq_description") or task.get("boq_name") or "")
    return "待确认"


def pick_spec_from_mentions(material: dict[str, Any], keywords: list[str]) -> str:
    mentions = material.get("evidence_material_mentions") or []
    values = []
    for mention in mentions:
        raw = str(mention.get("raw_text") or "")
        spec = str(mention.get("spec") or "")
        if not any(keyword in raw or keyword in spec for keyword in keywords):
            continue
        if spec:
            values.append(spec)
    return unique_join(values)


def build_technical_param(material: dict[str, Any], task: dict[str, Any], spec: str) -> str:
    name = material.get("standard_name", "")
    desc = str(task.get("boq_description") or "")
    if name in ["防滑地砖", "墙面砖"] and spec != "待确认":
        return f"规格{spec}"
    if name == "腻子" and spec != "待确认":
        return f"腻子{spec}"
    if name == "乳胶漆" and spec != "待确认":
        return spec
    if name == "石膏找平材料" and spec != "待确认":
        return spec
    if name == "成品栏杆" and spec != "待确认":
        return spec
    if name == "成品门窗":
        return desc or spec
    return "待确认"


def infer_plan_quantity(task: dict[str, Any], material: dict[str, Any]) -> tuple[Any, str, str]:
    boq_qty = task.get("boq_quantity")
    boq_unit = task.get("boq_unit")
    name = material.get("standard_name", "")

    if boq_qty in [None, ""]:
        return "", material.get("purchase_unit", ""), "待确认"

    if name in ["防滑地砖", "墙面砖", "纸面石膏板"]:
        return boq_qty, "m²", "按 BOQ 工程量暂列，损耗和规格换算待确认"
    if name in ["成品栏杆"]:
        return boq_qty, "m", "按 BOQ 工程量暂列，节点和材质待确认"
    if name in ["成品门窗"]:
        return boq_qty, boq_unit or material.get("purchase_unit", ""), "按 BOQ 工程量暂列，需门窗表确认"
    if name in ["腻子", "乳胶漆", "水泥砂浆", "美缝剂", "轻钢龙骨", "石膏找平材料"]:
        return "", material.get("purchase_unit", ""), "需补理论用量/厚度/间距/损耗率后计算"

    return "", material.get("purchase_unit", ""), "待确认"


def material_master_by_id(project: str) -> dict[str, dict[str, Any]]:
    path = PROJECTS_DIR / project / "材料拆解工作区" / "materials_master.json"
    if not path.exists():
        return {}
    return {item["material_id"]: item for item in read_json(path)}


def with_source_confirmed_supplements(project: str, task: dict[str, Any], materials: list[dict[str, Any]]) -> list[dict[str, Any]]:
    desc = f"{task.get('boq_name', '')} {task.get('boq_description', '')}"
    existing = {m.get("standard_name") for m in materials}
    master = material_master_by_id(project)
    supplements = []
    if "石膏" in desc and "石膏找平材料" not in existing:
        base = dict(master.get("MAT-003", {}))
        if base:
            supplements.append({
                "line_id": f"{task['boq_item_id']}-M-SG",
                "material_id": "MAT-003",
                "standard_name": base["standard_name"],
                "role": "配套材料",
                "purchase_unit": base["purchase_unit"],
                "required_specs": base.get("required_specs", []),
                "evidence_material_mentions": [],
            })
    return [*materials, *supplements]


def build_rows(project: str) -> list[dict[str, Any]]:
    root = PROJECTS_DIR / project
    workspace_dir = root / "材料拆解工作区"
    tasks = read_json(workspace_dir / "拆解任务队列.json")

    rows = []
    seq = 1
    for task in tasks:
        task_materials = with_source_confirmed_supplements(project, task, task.get("proposed_materials", []))
        for material in task_materials:
            plan_qty, plan_unit, remark = infer_plan_quantity(task, material)
            spec = pick_spec(material, task)
            technical_param = build_technical_param(material, task, spec)
            rows.append({
                "序号": seq,
                "项目名称": project,
                "清单编号": task.get("boq_item_id"),
                "清单名称": task.get("boq_name"),
                "材料名称": material.get("standard_name"),
                "材料角色": material.get("role"),
                "规格型号": spec,
                "技术参数": technical_param,
                "工程量": task.get("boq_quantity"),
                "工程量单位": task.get("boq_unit"),
                "计划采购量": plan_qty,
                "采购单位": plan_unit,
                "损耗率": "待确认",
                "采购状态": "待确认" if not plan_qty else "暂列",
                "备注": remark,
            })
            seq += 1
    return rows


def write_workbook(path: Path, rows: list[dict[str, Any]]) -> None:
    headers = [
        "序号", "项目名称", "清单编号", "清单名称", "材料名称", "材料角色",
        "规格型号", "技术参数", "工程量", "工程量单位", "计划采购量",
        "采购单位", "损耗率", "采购状态", "备注",
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "采购计划物料清单"
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = {
        "A": 8, "B": 14, "C": 12, "D": 26, "E": 18, "F": 14,
        "G": 26, "H": 38, "I": 12, "J": 12, "K": 14, "L": 12,
        "M": 12, "N": 12, "O": 34,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(idx)].bestFit = False

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", default="宿州302")
    args = parser.parse_args()

    root = PROJECTS_DIR / args.project
    if not root.exists():
        raise FileNotFoundError(f"项目不存在: {root}")
    if not (root / "材料拆解工作区" / "拆解任务队列.json").exists():
        raise FileNotFoundError("缺少材料拆解工作区，请先运行 python3 运行入口.py --stage workspace")

    rows = build_rows(args.project)
    output_path = root / "最终输出" / "采购计划物料清单.xlsx"
    write_workbook(output_path, rows)

    manifest = {
        "project": args.project,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "output_file": str(output_path),
        "rows": len(rows),
    }
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
