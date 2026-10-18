# -*- coding: utf-8 -*-
"""客户全量信息与企业账号全景导出引擎 (Enterprise Admin Export Engine)。

麦肯锡咨询风格专业多 Sheet Excel 报表：
- Sheet 1: 客户预约与商务咨询台账 (Leads & Inquiries)
- Sheet 2: 租户企业与账号权益台账 (Tenants & Accounts)
- Sheet 3: AI视觉模型与BYOK密钥配置 (Model & Keys Hub)
"""
import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from db import db_list_leads, db_list_registrations
import store


def _apply_sheet_styling(ws, headers: list[str], rows: list[list]):
    """应用商业咨询级表格排版风格。"""
    # 颜色与边框定义
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    thin_side = Side(style="thin", color="E2E8F0")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # 1. 写入表头
    ws.append(headers)
    ws.row_dimensions[1].height = 28
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = thin_border

    # 2. 写入数据行
    for r_idx, row_data in enumerate(rows, start=2):
        ws.append(row_data)
        ws.row_dimensions[r_idx].height = 22
        fill = zebra_fill if r_idx % 2 == 0 else white_fill
        for c_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.fill = fill
            cell.border = thin_border
            cell.font = Font(name="Microsoft YaHei", size=10)
            # 文本较长居左，状态/日期/编号/数字居中
            val_str = str(cell.value or "")
            if len(val_str) > 20 or "\n" in val_str or any(k in headers[c_idx - 1] for k in ["备注", "说明", "留言", "地址", "公司", "企业"]):
                cell.alignment = align_left
            else:
                cell.alignment = align_center

    # 3. 自动计算列宽
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            # 中文字符算 2 个宽度，英文字符算 1 个
            cell_len = sum(2 if ord(c) > 127 else 1 for c in val)
            if cell_len > max_len:
                max_len = cell_len
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 48)

    # 4. 冻结首行与开启筛选
    ws.freeze_panes = "A2"
    if ws.dimensions:
        ws.auto_filter.ref = ws.dimensions


def generate_admin_export_wb() -> openpyxl.Workbook:
    """生成全量导出的多 Sheet Workbook。"""
    wb = openpyxl.Workbook()
    # 默认工作表改为 Sheet 1
    ws_leads = wb.active
    ws_leads.title = "客户预约与咨询台账"
    ws_users = wb.create_sheet(title="租户企业与账号权益")
    ws_keys = wb.create_sheet(title="AI模型与BYOK密钥配置")

    # ---------- Sheet 1: 客户预约与咨询台账 ----------
    leads = db_list_leads()
    lead_headers = [
        "序号", "客户姓名", "联系电话", "所属公司/机构",
        "需求内容与留言", "业务来源", "跟进状态", "商务跟进备注",
        "提交时间", "最新跟进时间", "线索ID"
    ]
    lead_rows = []
    status_map = {
        "pending": "待跟进",
        "contacted": "已联系",
        "converted": "已签约转化",
        "rejected": "暂无意向",
        "invalid": "无效信息",
    }
    for idx, ld in enumerate(leads, start=1):
        st = status_map.get(ld.get("status", ""), ld.get("status", "待跟进"))
        lead_rows.append([
            idx,
            ld.get("name", ""),
            ld.get("phone", ""),
            ld.get("company", ""),
            ld.get("content", ""),
            ld.get("source", "landing_priority_access"),
            st,
            ld.get("notes", ""),
            ld.get("created_at", ""),
            ld.get("updated_at", ld.get("created_at", "")),
            ld.get("id", ""),
        ])
    _apply_sheet_styling(ws_leads, lead_headers, lead_rows)

    # ---------- Sheet 2: 租户企业与账号权益 ----------
    users = db_list_registrations()
    user_headers = [
        "序号", "登录账号", "姓名/昵称", "角色权限", "账号状态",
        "联系方式", "所属企业名称", "企业ID", "企业状态",
        "名下图纸任务数", "名下工程项目数", "注册时间", "最近登录时间", "用户ID"
    ]
    user_rows = []
    user_status_map = {
        "active": "正常生效",
        "frozen": "已冻结停用",
    }
    for idx, u in enumerate(users, start=1):
        ust = user_status_map.get(u.get("status", "active"), u.get("status", "正常生效"))
        tst = "正常营业" if u.get("tenant_status") == "active" else "暂停"
        role_name = "租户管理员" if u.get("role") == "admin" else ("超级管理员" if u.get("role") == "super_admin" else "普通成员")
        user_rows.append([
            idx,
            u.get("username", ""),
            u.get("display_name", ""),
            role_name,
            ust,
            u.get("contact", ""),
            u.get("tenant_name", ""),
            u.get("tenant_id", ""),
            tst,
            u.get("jobs_count", 0),
            u.get("projects_count", 0),
            u.get("created_at", ""),
            u.get("last_login_at", ""),
            u.get("user_id", ""),
        ])
    _apply_sheet_styling(ws_users, user_headers, user_rows)

    # ---------- Sheet 3: AI模型与BYOK密钥配置 ----------
    configs = store.list_all_tenant_model_configs(mask=True)
    key_headers = [
        "序号", "企业名称", "企业租户ID", "配置归属", "企业状态",
        "视觉服务商", "模型代号", "API接口地址", "Key配置状态", "API Key (安全脱敏)"
    ]
    key_rows = []
    for idx, cfg in enumerate(configs, start=1):
        key_rows.append([
            idx,
            cfg.get("tenant_name", ""),
            cfg.get("tenant_id", ""),
            "系统全局默认" if cfg.get("is_system_default") else "租户独立BYOK",
            "正常" if cfg.get("tenant_status") == "active" else "停用",
            cfg.get("vision_provider", "qwen"),
            cfg.get("vision_model", ""),
            cfg.get("vision_base_url", ""),
            "已配置有效" if cfg.get("vision_configured") else "未配置(无Key)",
            cfg.get("api_key_masked", ""),
        ])
    _apply_sheet_styling(ws_keys, key_headers, key_rows)

    return wb


def generate_admin_export_bytes() -> bytes:
    """生成全量导出的 Excel 字节流（可直接用于 HTTP 响应下载）。"""
    wb = generate_admin_export_wb()
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
