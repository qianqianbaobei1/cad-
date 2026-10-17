# -*- coding: utf-8 -*-
"""图纸原生读取（native_takeoff）回归：各类对象不丢字、图块计数、线长、表格、导出。"""
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import ezdxf
import openpyxl
import pymupdf

from extractor.native_takeoff import (
    build_native_workbook, inventory_from_json, scan_cad, scan_pdf, supplement_cad_texts,
)


def _make_dxf(path: str) -> str:
    doc = ezdxf.new("R2018", setup=True)
    msp = doc.modelspace()
    msp.add_text("1AL1 照明配电箱", dxfattribs={"height": 5, "layer": "TEXT"}).set_placement((0, 0))

    # 块：内部文字 + 一条线 + 属性定义
    blk = doc.blocks.new("LAMP")
    blk.add_line((0, 0), (10, 0))
    blk.add_text("灯", dxfattribs={"height": 2}).set_placement((1, 1))
    blk.add_attdef("NAME", (0, -3), dxfattribs={"height": 2})
    for i in range(3):
        ins = msp.add_blockref("LAMP", (100 + i * 20, 0), dxfattribs={"layer": "EQUIP"})
        ins.add_attrib("NAME", "双管荧光灯", (100 + i * 20, -3))

    # 阵列块 MINSERT 2 行 × 3 列 = 6 个
    m = msp.add_blockref("LAMP", (0, 200), dxfattribs={"layer": "EQUIP"})
    m.dxf.row_count, m.dxf.column_count = 2, 3
    m.dxf.row_spacing, m.dxf.column_spacing = 50, 50

    # 嵌套块：外层块里放 LAMP
    outer = doc.blocks.new("PANEL")
    outer.add_mtext("嵌套文字", dxfattribs={"char_height": 3}).set_location((0, 20))
    outer.add_blockref("LAMP", (0, 0))
    msp.add_blockref("PANEL", (500, 500))

    # 尺寸标注（显示文字在标注图形块里）
    dim = msp.add_linear_dim(base=(0, -50), p1=(0, -40), p2=(1000, -40), dimstyle="EZDXF")
    dim.render()

    # 电缆线：图层线长 1000 + 500
    msp.add_line((0, -100), (1000, -100), dxfattribs={"layer": "WIRE"})
    msp.add_lwpolyline([(0, -200), (500, -200)], dxfattribs={"layer": "WIRE"})

    # 字体缺失的文字样式
    doc.styles.add("缺字体", font="not_exist_font.shx")
    msp.add_text("缺字体文字", dxfattribs={"height": 3, "style": "缺字体"}).set_placement((0, 50))
    doc.saveas(path)
    return path


class NativeTakeoffTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.dxf = _make_dxf(os.path.join(cls.tmp, "sample.dxf"))
        cls.inv = scan_cad(cls.dxf)

    def texts(self, kind=None):
        return [t.text for t in self.inv.texts if kind is None or t.kind == kind]

    def test_reads_all_text_kinds(self):
        all_texts = self.texts()
        self.assertIn("1AL1 照明配电箱", all_texts)
        self.assertIn("嵌套文字", all_texts)             # 块内多行文字
        self.assertGreaterEqual(all_texts.count("灯"), 3 + 6 + 1)  # 普通块 + 阵列 + 嵌套
        self.assertEqual(self.texts("ATTRIB").count("双管荧光灯"), 3)
        self.assertTrue(any("1000" in t for t in self.texts("DIMENSION")))

    def test_block_counts_include_minsert(self):
        lamp = next(b for b in self.inv.blocks if b["name"] == "LAMP")
        self.assertEqual(lamp["count"], 3 + 6)          # 嵌套在 PANEL 里的不计入顶层统计
        self.assertIn("双管荧光灯", lamp["attributes"])
        self.assertIn("灯", lamp["sample_text"])

    def test_layer_length(self):
        wire = next(r for r in self.inv.layer_lengths if r["layer"] == "WIRE")
        self.assertAlmostEqual(wire["length"], 1500.0, places=1)
        self.assertEqual(wire["count"], 2)

    def test_font_report_flags_missing(self):
        row = next(f for f in self.inv.fonts if f["style"] == "缺字体")
        self.assertFalse(row["font_found"])
        self.assertEqual(row["used"], 1)
        self.assertTrue(any("not_exist_font" in w for w in self.inv.warnings))

    def test_supplement_dedupes_existing(self):
        from extractor.cad import extract_cad_entities
        base = extract_cad_entities(self.dxf)
        merged = supplement_cad_texts(self.dxf, base, self.inv)
        attribs = [t for t in merged if t["text"] == "双管荧光灯"]
        self.assertEqual(len(attribs), 3)                # 顶层属性不重复追加
        self.assertIn("嵌套文字", [t["text"] for t in merged])
        self.assertGreater(len(merged), len(base))

    def test_workbook_export_roundtrip(self):
        data = self.inv.to_json()
        out = os.path.join(self.tmp, "native.xlsx")
        build_native_workbook(inventory_from_json(data), out)
        wb = openpyxl.load_workbook(out, read_only=True)
        self.assertEqual([ws.title for ws in wb.worksheets],
                         ["读取概览", "文字清单", "表格", "图块统计", "图层线长", "字体检查"])
        self.assertEqual(wb["文字清单"].max_row, len(self.inv.texts) + 1)


class PdfNativeTests(unittest.TestCase):
    def test_pdf_text_layer_and_grid_table(self):
        tmp = tempfile.mkdtemp()
        path = os.path.join(tmp, "t.pdf")
        doc = pymupdf.open()
        page = doc.new_page(width=600, height=400)
        xs, ys = [50, 200, 350], [50, 100, 150, 200]
        for y in ys:
            page.draw_line((xs[0], y), (xs[-1], y))
        for x in xs:
            page.draw_line((x, ys[0]), (x, ys[-1]))
        cells = [["回路", "断路器"], ["WL1", "C16/1P"], ["WL2", "C20/1P"]]
        for r, row in enumerate(cells):
            for c, val in enumerate(row):
                page.insert_text((xs[c] + 10, ys[r] + 30), val, fontname="china-s", fontsize=11)
        doc.save(path)
        inv = scan_pdf(path, use_ocr=False)
        self.assertIn("C16/1P", [t.text for t in inv.texts])
        self.assertEqual(inv.texts[0].kind, "PDF_TEXT")
        self.assertEqual(len(inv.tables), 1)
        flat = [c for row in inv.tables[0]["cells"] for c in row]
        for val in ("WL1", "C16/1P", "C20/1P"):
            self.assertIn(val, flat)


def _tangent_cipher(plain: str) -> str:
    """天正默认分支（key=0x528, base='A'）加密：每字节 ^0x05 后拆成两位 26 进制大写字母。"""
    out = []
    for b in plain.encode("gb18030"):
        t = b ^ 0x05
        out.append(chr(0x41 + t // 26) + chr(0x41 + t % 26))
    return "".join(out)


class TangentDecryptTests(unittest.TestCase):
    """天正块（$EQUIP$…）里的加密文字与属性值解出明文；普通块与普通大写标注不动。"""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        path = os.path.join(cls.tmp, "tangent.dxf")
        doc = ezdxf.new("R2018")
        msp = doc.modelspace()
        blk = doc.blocks.new("$EQUIP$00001")
        blk.add_text(_tangent_cipher("照明配电箱"), dxfattribs={"height": 2}).set_placement((0, 0))
        blk.add_text("WDZN", dxfattribs={"height": 2}).set_placement((0, 5))
        blk.add_attdef("NAME", (0, -3), dxfattribs={"height": 2})
        ins = msp.add_blockref("$EQUIP$00001", (0, 0))
        ins.add_attrib("NAME", _tangent_cipher("插座"), (0, -3))
        plain = doc.blocks.new("PLAIN")
        plain.add_text(_tangent_cipher("照明"), dxfattribs={"height": 2}).set_placement((0, 0))
        msp.add_blockref("PLAIN", (100, 0))
        doc.saveas(path)
        cls.inv = scan_cad(path)

    def test_tangent_block_texts_decrypted(self):
        texts = [t.text for t in self.inv.texts]
        self.assertIn("照明配电箱", texts)
        self.assertIn("插座", texts)
        self.assertIn("WDZN", texts)                      # 普通大写标注不被误解
        self.assertIn(_tangent_cipher("照明"), texts)     # 非天正块不尝试解密
        self.assertTrue(any("天正加密文字" in w for w in self.inv.warnings))

    def test_block_statistics_decrypted(self):
        row = next(b for b in self.inv.blocks if b["name"] == "$EQUIP$00001")
        self.assertTrue(row["tangent"])
        self.assertIn("照明配电箱", row["sample_text"])
        self.assertIn("插座", row["attributes"])


class ExtractionSupplementTests(unittest.TestCase):
    """extraction_supplements：给箱柜/回路识别补现有读取漏掉的文字。"""

    def test_minsert_copies_and_mleader(self):
        from extractor.native_takeoff import extraction_supplements
        doc = ezdxf.new("R2018", setup=True)
        msp = doc.modelspace()
        blk = doc.blocks.new("TAG")
        blk.add_text("WL1", dxfattribs={"height": 2}).set_placement((0, 0))
        m = msp.add_blockref("TAG", (0, 0))
        m.dxf.row_count, m.dxf.column_count = 1, 4
        m.dxf.column_spacing = 100
        ml = msp.add_multileader_mtext("Standard")
        ml.set_content("C65N-C16/1P", char_height=2.5)
        ml.add_leader_line(ezdxf.render.mleader.ConnectionSide.left, [ezdxf.math.Vec2(50, 50)])
        ml.build(insert=ezdxf.math.Vec2(60, 60))
        texts, tables = extraction_supplements(doc)
        wl = sorted(round(t["x"], -2) for t in texts if t["text"] == "WL1")   # x 为文字中心
        self.assertEqual(wl, [100, 200, 300])            # 第一个副本由现有读取负责，这里补其余 3 个
        self.assertIn("C65N-C16/1P", [t["text"] for t in texts])
        self.assertEqual(tables, [])


class NativeBoxInventoryTests(unittest.TestCase):
    """native_box_inventory：系统图标题 / 平面图例属性 / “箱号:回路”标注列出配电箱。"""

    def test_levels_and_filters(self):
        from extractor.native_takeoff import native_box_inventory
        items = [
            ("02AP01 配电总箱1", "TEXT"),                      # 系统图标题 → confirmed
            ("02AL01 照明配电箱", "TEXT"),
            ("02AL01:N1", "TEXT"), ("02AL01:N2", "TEXT"), ("02AL01：N2", "TEXT"),
            ("02AL01", "ATTRIB"),                               # 平面图例属性
            ("02PB1:N1", "TEXT"), ("02PB1", "ATTRIB"),          # 只在平面图 → plan
            ("40AH 直流屏组", "TEXT"),                          # 额定值，不是箱号
            ("ET1D07 低压柜中包含两路电梯供电回路，此柜需要三个切非消信号.", "TEXT"),  # 说明句子
            ("T11＃变压器柜", "TEXT"),                          # “#”序号写法
        ]
        inv = native_box_inventory(items)
        self.assertEqual(inv["02AP01"]["level"], "confirmed")
        self.assertEqual(inv["02AP01"]["name"], "配电总箱1")
        self.assertEqual(inv["02AL01"]["plan_circuits"], ["N1", "N2"])
        self.assertEqual(inv["02AL01"]["symbols"], 1)
        self.assertEqual(inv["02PB1"]["level"], "plan")
        for bad in ("40AH", "ET1D07", "T11"):
            self.assertNotEqual(inv.get(bad, {}).get("level"), "confirmed", bad)


if __name__ == "__main__":
    unittest.main()
