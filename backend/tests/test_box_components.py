# -*- coding: utf-8 -*-
"""T17/T13 回归：元器件按柜拆分同源，控制电缆回路不进断路器组价。

T17：data.components 是项目级汇总，整行 quantity 是全项目总数。旧组价只判断
"used_in 里有没有本柜"就把整行交给本柜，锅炉房真实任务因此把 232 件算成 757 件。
T13：ZR-KYJVP 控制电缆回路（至现场操作柱/设备PLC柜）breaker 为空，不能被补成
"分支断路器"，也不该计入断路器回路数与出线组装工时。
"""
import json
import os
import unittest
import uuid

import test_support  # noqa: F401
from fastapi.testclient import TestClient

from app import app as fastapi_app, jobs
from db import db_save_job
from extractor.box_components import (
    UNASSIGNED,
    component_total,
    split_components_by_box,
)
from extractor.pricing import calculate_box_quotation

FIXTURE = os.path.join(os.path.dirname(__file__), "fixtures", "job_b1d7db50ac00_data.json")


def _snapshot() -> dict:
    with open(FIXTURE, encoding="utf-8") as fh:
        return json.load(fh)


def _mk_job(data: dict) -> str:
    job_id = "t17_" + uuid.uuid4().hex[:8]
    jobs[job_id] = {
        "job_id": job_id,
        "filename": "按柜拆分回归.dwg",
        "file_type": "dwg",
        "status": "done",
        "project": "回归项目",
        "summary": {"title": "按柜拆分回归"},
        "data": data,
        "changes": [],
    }
    db_save_job(jobs[job_id])
    return job_id


def _box_rows(rows_by_box: dict, code: str) -> list[dict]:
    return rows_by_box.get(code) or []


def _quantities(rows: list[dict]) -> dict[tuple[str, str], int]:
    out: dict[tuple[str, str], int] = {}
    for row in rows:
        key = (row["name"], row["spec"])
        out[key] = out.get(key, 0) + int(row["quantity"])
    return out


class TestSplitComponentsByBox(unittest.TestCase):
    def test_跨柜一行按片段数拆开且总量守恒(self):
        data = {
            "boxes": [{"code": "1AP1"}, {"code": "1AP2"}, {"code": "1AP3"}],
            "circuits": [],
            "components": [
                {"name": "塑壳断路器", "spec": "MCCB-100/50A/3P", "unit": "只", "quantity": 6,
                 "used_in": "1AP1 D01、1AP1 D02、1AP2 D01、1AP2 D02、1AP3 D01、1AP3 D02"},
            ],
        }
        split = split_components_by_box(data)
        for code in ("1AP1", "1AP2", "1AP3"):
            self.assertEqual([row["quantity"] for row in split[code]], [2], code)
        self.assertEqual(component_total(split), 6)

    def test_十三个柜各一件(self):
        codes = [f"0{i}AL" for i in range(1, 14)]
        data = {
            "boxes": [{"code": code} for code in codes],
            "circuits": [],
            "components": [
                {"name": "SPD后备保护器", "spec": "SPD厂家配套", "unit": "只", "quantity": 13,
                 "used_in": "、".join(codes)},
            ],
        }
        split = split_components_by_box(data)
        for code in codes:
            self.assertEqual([row["quantity"] for row in split[code]], [1], code)
        self.assertEqual(component_total(split), 13)

    def test_归不到任何柜的数量不摊给别人_单列未归属并标存疑(self):
        data = {
            "boxes": [{"code": "1AP1"}, {"code": "1AP2"}],
            "circuits": [],
            "components": [
                {"name": "微型断路器", "spec": "MCB-C16A/1P", "unit": "只", "quantity": 2,
                 "used_in": "1AP1 D01、1AP2 D01"},
                {"name": "多功能仪表", "spec": "带通讯接口多功能表", "unit": "只", "quantity": 3,
                 "used_in": "图纸未标柜号"},
            ],
        }
        split = split_components_by_box(data)
        self.assertFalse([r for r in split["1AP1"] if r["name"] == "多功能仪表"])
        self.assertFalse([r for r in split["1AP2"] if r["name"] == "多功能仪表"])
        orphan = _box_rows(split, UNASSIGNED)
        self.assertEqual(len(orphan), 1)
        self.assertEqual(orphan[0]["quantity"], 3)
        self.assertIn("存疑", orphan[0]["note"])
        # 未归属也留在总数里，一件都不许丢
        self.assertEqual(component_total(split), 5)

    def test_片段数与数量不等时按回路逐条明细落到各柜(self):
        """02AL01 8 条回路、02AL02 2 条回路，quantity=10：不能按 8 段比例算成 9/1。"""
        data = {
            "boxes": [{"code": "02AL01"}, {"code": "02AL02"}],
            "circuits": [
                {"box": "02AL01", "circuit_no": f"N{i}", "breaker": "RCBO/2P C20"} for i in range(8)
            ] + [
                {"box": "02AL02", "circuit_no": f"N{i}", "breaker": "RCBO/2P C20"} for i in range(2)
            ],
            "components": [
                # used_in 里同柜的"备用"回路重名，只留下 7 + 1 段，但真实件数是 10
                {"name": "剩余电流动作断路器", "spec": "RCBO/2P C20", "unit": "只", "quantity": 10,
                 "used_in": "02AL01 N0、02AL01 N1、02AL01 N2、02AL01 N3、02AL01 N4、02AL01 N5、"
                            "02AL01 N6、02AL02 N0"},
            ],
        }
        split = split_components_by_box(data)
        self.assertEqual(split["02AL01"][0]["quantity"], 8)
        self.assertEqual(split["02AL02"][0]["quantity"], 2)

    def test_真实任务快照_各柜合计等于项目合计232件(self):
        data = _snapshot()
        split = split_components_by_box(data)
        expected = sum(float(c["quantity"]) for c in data["components"] if c["name"] != "配电箱体")
        self.assertEqual(expected, 232)
        self.assertEqual(component_total(split), 232)
        self.assertFalse(_box_rows(split, UNASSIGNED))
        # 哪一柜拿到的件数都不能超过该规格的全项目数量（旧逻辑把整行给每个命中的柜）
        project_qty = {(c["name"], c["spec"]): int(c["quantity"]) for c in data["components"]}
        for code, rows in split.items():
            for row in rows:
                self.assertLessEqual(row["quantity"], project_qty[(row["name"], row["spec"])], code)
        # 02AP1 的 MCCB 是 2 只（拆量页原来就显示 2，组价不能再算 6）
        mccb = [r for r in split["02AP1"] if r["spec"] == "MCCB-160MA/50A/3P"]
        self.assertEqual([r["quantity"] for r in mccb], [2])

    def test_真实任务快照_组价接口各柜数量与拆分函数一致(self):
        data = _snapshot()
        split = split_components_by_box(data)
        client = TestClient(fastapi_app)
        job_id = _mk_job(data)
        resp = client.get(f"/api/jobs/{job_id}/quotation")
        self.assertEqual(resp.status_code, 200)
        quote = resp.json()
        self.assertTrue(quote["boxes"], "组价结果不能没有箱柜")
        for box in quote["boxes"]:
            code = box["box_code"]
            quoted = {}
            for row in quote["rows"]:
                if row.get("type") != "component" or row.get("box_code") != code:
                    continue
                key = (row["name"], row["design_spec"])
                quoted[key] = quoted.get(key, 0) + int(row["quantity"] or 0)
            self.assertEqual(quoted, _quantities(_box_rows(split, code)), code)
        # 全项目组价元器件件数 = 232，不再是 757
        total_quoted = sum(
            int(row["quantity"] or 0)
            for row in quote["rows"] if row.get("type") == "component"
        )
        self.assertEqual(total_quoted, 232)

    def test_拆量接口与组价同源(self):
        """拆量页每柜元器件直接吃后端拆分结果，与组价页同名同规格同数量。"""
        data = _snapshot()
        client = TestClient(fastapi_app)
        job_id = _mk_job(data)
        detail = client.get(f"/api/jobs/{job_id}").json()
        split = split_components_by_box(data)
        served = {code: _quantities(rows) for code, rows in detail["data"]["box_components"].items() if rows}
        self.assertEqual(served, {code: _quantities(rows) for code, rows in split.items() if rows})
        self.assertEqual(_quantities(detail["data"]["box_components"]["02AP1"]),
                         _quantities(split["02AP1"]))
        # 派生字段只挂在响应上，不写回任务缓存/任务数据
        self.assertNotIn("box_components", jobs[job_id]["data"])


class TestControlCableCircuits(unittest.TestCase):
    def _box(self):
        return [{"code": "02AP1", "size": "GGD", "ip_rating": "IP54"}]

    def _split(self, circuits):
        """断路器行由 box_components 按回路补出：组价不再自己合成。"""
        box = self._box()[0]
        return split_components_by_box({"boxes": [box], "circuits": circuits, "components": []}).get("02AP1", [])

    def test_控制电缆与操作柱回路不生成断路器也不计回路工时(self):
        circuits = [
            {"box": "02AP1", "circuit_no": "D01", "breaker": "MCCB-160MA/50A/3P", "load_name": "1#锅炉给水泵"},
            {"box": "02AP1", "circuit_no": "D01-C1", "cable": "ZR-KYJVP-0.6/1KV-10x1.5-SC32-CT",
             "load_name": "至现场操作柱"},
            {"box": "02AP1", "circuit_no": "D01-C2", "cable": "ZR-KYJVP-0.6/1KV-7x1.5-SC32-CT",
             "load_name": "至设备PLC柜"},
        ]
        q = calculate_box_quotation(self._box()[0], circuits, self._split(circuits), brand="正泰")
        self.assertEqual(len(q["components"]), 1)
        self.assertEqual(q["components"][0]["spec"], "MCCB-160MA/50A/3P")
        # 控制电缆回路没有被补成“分支断路器（规格未标注）”
        self.assertTrue(all(row["spec"] for row in q["components"]))
        # 只有 1 条断路器回路：壳体按 1 回路取档，工时只算 1 条出线
        self.assertEqual(q["circuits_count"], 1)
        self.assertEqual(q["control_circuits_count"], 2)
        self.assertEqual(q["cost_breakdown"]["labor_cost"], 45.0)

    def test_只有控制电缆回路的箱柜不出任何断路器(self):
        circuits = [
            {"box": "02AP1", "circuit_no": "D01-C1", "cable": "ZR-KYJVP-0.6/1KV-10x1.5-SC32-CT",
             "load_name": "至现场操作柱"},
            {"box": "02AP1", "circuit_no": "D01-C2", "cable": "ZR-KYJVP-0.6/1KV-7x1.5-SC32-CT",
             "load_name": "至设备PLC柜"},
        ]
        q = calculate_box_quotation(self._box()[0], circuits, self._split(circuits), brand="正泰")
        self.assertEqual(q["components"], [])
        self.assertEqual(q["circuits_count"], 0)
        self.assertEqual(q["control_circuits_count"], 2)

    def test_控制电缆回路仍留在回路清单里(self):
        data = _snapshot()
        control = [
            c for c in data["circuits"]
            if c.get("box") == "02AP1" and str(c.get("cable") or "").upper().find("KYJVP") >= 0
        ]
        self.assertEqual(len(control), 6)
        split = split_components_by_box(data)
        # 控制电缆回路不产生任何元器件行
        for row in split["02AP1"]:
            self.assertNotIn("至现场操作柱", row["spec"])
            self.assertNotIn("KYJVP", row["spec"])
            self.assertNotIn("分支断路器", row["name"])


class TestSynthesizedBreakers(unittest.TestCase):
    """T17b：断路器只在一处合成——没有断路器元器件的柜在拆分时按回路补行，组价不再自算。"""

    def _data(self):
        return {
            "boxes": [{"code": "03AP1", "name": "动力箱", "size": "800×2200×600"}],
            "circuits": [
                {"box": "03AP1", "circuit_no": "进线", "load_name": "市电进线", "breaker": "NM1-125S/3300 100A"},
                {"box": "03AP1", "circuit_no": "D01", "load_name": "风机", "breaker": "MCCB-160MA/50A/3P"},
                {"box": "03AP1", "circuit_no": "D02", "load_name": "至现场操作柱",
                 "cable": "ZR-KYJVP-0.6/1KV-10x1.5-SC32-CT"},
            ],
            # 元器件汇总里一个断路器都没有（旧逻辑只有组价会自己合成）
            "components": [{"name": "浪涌保护器", "spec": "SPD 4P", "unit": "只", "quantity": 1,
                            "used_in": "03AP1"}],
        }

    def test_拆量为没有断路器的柜补出断路器行(self):
        split = split_components_by_box(self._data())
        rows = split["03AP1"]
        breakers = [(r["name"], r["spec"], r["quantity"]) for r in rows if "断路器" in r["name"]]
        self.assertEqual(breakers, [("进线断路器", "NM1-125S/3300 100A", 1),
                                    ("分支断路器", "MCCB-160MA/50A/3P", 1)])
        # 补出来的行标清来源，控制电缆回路不补
        for row in rows:
            if "断路器" in row["name"]:
                self.assertIn("由回路图提取", row["note"])
        self.assertEqual(component_total(split), 3)

    def test_拆量与组价同一份断路器清单(self):
        data = self._data()
        split = split_components_by_box(data)
        q = calculate_box_quotation(data["boxes"][0], data["circuits"], split["03AP1"], brand="正泰")
        quoted = sorted((row["name"], row["spec"], row["quantity"]) for row in q["components"])
        expected = sorted((row["name"], row["spec"], row["quantity"]) for row in split["03AP1"])
        self.assertEqual(quoted, expected)


class TestMissingQuantity(unittest.TestCase):
    """T17b：数量缺失/0 不得凭空记 1 件，记一条“数量未标注”待核对。"""

    def test_零数量不计件且进待核对(self):
        from extractor.box_components import _as_count, pending_checks

        self.assertEqual(_as_count(0), 0)
        self.assertEqual(_as_count(None), 0)
        self.assertEqual(_as_count(""), 0)
        self.assertEqual(_as_count(3.0), 3)

        data = {
            "boxes": [{"code": "04AP1"}],
            "circuits": [],
            "components": [{"name": "电压表", "spec": "电压表", "unit": "只", "quantity": 0,
                            "used_in": "04AP1"}],
        }
        split = split_components_by_box(data)
        self.assertEqual(component_total(split), 0)
        checks = pending_checks(split)
        self.assertEqual(len(checks), 1)
        self.assertIn("数量未标注", checks[0]["note"])
        self.assertEqual(checks[0]["kind"], "doubt")


class TestPendingChecks(unittest.TestCase):
    """T17b：未归属/存疑行作为派生待核对清单随接口返回，不落盘。"""

    def test_未归属行进入派生待核对清单且不写回任务(self):
        from extractor.box_components import pending_checks

        data = {
            "boxes": [{"code": "05AP1"}, {"code": "05AP2"}],
            "circuits": [],
            "components": [
                {"name": "多功能仪表", "spec": "带通讯接口多功能表", "unit": "只", "quantity": 2,
                 "used_in": "图纸未标柜号"},
                {"name": "微型断路器", "spec": "MCB-C16A/1P", "unit": "只", "quantity": 1,
                 "used_in": "05AP1 N1"},
            ],
        }
        checks = pending_checks(split_components_by_box(data))
        self.assertEqual(len(checks), 1)
        self.assertEqual(checks[0]["kind"], "unassigned")
        self.assertEqual(checks[0]["box_code"], "")
        self.assertEqual(checks[0]["quantity"], 2)

        # 接口返回派生字段，任务缓存/任务数据里没有
        job_id = _mk_job(data)
        served = TestClient(fastapi_app).get(f"/api/jobs/{job_id}").json()
        self.assertEqual(served["data"]["pending_checks"], checks)
        self.assertNotIn("pending_checks", jobs[job_id]["data"])

    def test_真实快照没有待核对行时不编造(self):
        from extractor.box_components import pending_checks

        self.assertEqual(pending_checks(split_components_by_box(_snapshot())), [])


class TestHeaderNoiseRows(unittest.TestCase):
    """T17b：表头/字段名被当成元器件时在 assemble 源头过滤，并写进待核对（INFO）。"""

    def _assemble(self, devices):
        from extractor.schema import Box, Circuit, RawExtraction
        from extractor.assemble import assemble

        raw = RawExtraction(
            boxes=[Box(code="C01", name="低压配电柜 C01", ip_rating="IP30")],
            circuits=[Circuit(box="C01", circuit_no="E1", breaker="MCB-C16A/1P")],
            extra_devices=devices,
            requirements=[],
            uncertainties=[],
        )
        return assemble(raw)

    def test_表头文字不成元器件且留下INFO待核对(self):
        from extractor.schema import ExtraDevice

        result = self._assemble([
            ExtraDevice(name="回路名称", spec="机械应急启动柜", unit="只", quantity=1, used_in="C01"),
            ExtraDevice(name="低压配电柜型号 GGD", spec="GGD", unit="只", quantity=2, used_in="C01"),
            ExtraDevice(name="A\n计算电流", spec="371", unit="只", quantity=1, used_in="C01"),
            ExtraDevice(name="浪涌保护器", spec="SPD 4P", unit="只", quantity=1, used_in="C01"),
        ])
        names = [c.name for c in result.components]
        self.assertNotIn("回路名称", names)
        self.assertNotIn("低压配电柜型号 GGD", names)
        self.assertNotIn("A\n计算电流", names)
        self.assertIn("浪涌保护器", names)

        noise = [u for u in result.uncertainties if "疑似表头文字" in u.detail]
        self.assertEqual(len(noise), 3, [u.text for u in result.uncertainties])
        for item in noise:
            self.assertEqual(item.severity, "INFO")
            self.assertEqual(item.source, "program")


class TestTakeoffBomWorkbook(unittest.TestCase):
    """T17b：拆量导出吃按柜拆分结果，与拆量页同数同量。"""

    def _wb(self, data):
        import tempfile
        import openpyxl
        from extractor.excel import build_takeoff_bom_workbook

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
            path = tf.name
        build_takeoff_bom_workbook(
            {"job_id": "t17b", "filename": "锅炉房.dwg", "project": "锅炉房", "data": data}, path)
        self.addCleanup(os.unlink, path)
        return openpyxl.load_workbook(path, data_only=True)

    def _detail_rows(self, wb):
        ws = wb["箱柜及元器件明细"]
        return [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=ws.max_row)]

    def test_快照导出全项目232件且02AP1与拆量页一致(self):
        data = _snapshot()
        rows = self._detail_rows(self._wb(data))
        total = next(r[5] for r in rows if r[0] and "全项目工程量总计" in str(r[0]))
        self.assertEqual(total, 232)

        split = split_components_by_box(data)
        import re as _re
        mine = [r for r in rows if r[1] == "02AP1" and _re.fullmatch(r"\d+\.\d+", str(r[0] or ""))]
        self.assertEqual(len(mine), 5)
        self.assertEqual(sum(int(r[5]) for r in mine), 6)
        self.assertEqual({(r[2], r[3], int(r[5])) for r in mine},
                         {(r["name"], r["spec"], int(r["quantity"])) for r in split["02AP1"]})

    def test_导出里没有规格未标注的分支断路器且控制电缆只进回路明细(self):
        data = _snapshot()
        wb = self._wb(data)
        rows = self._detail_rows(wb)
        self.assertFalse([r for r in rows if r[3] and "规格未标注" in str(r[3])])
        self.assertFalse([r for r in rows if r[2] and "控制电缆出线" in str(r[2])])

        ws_c = wb["箱内回路明细"]
        circuit_rows = [[c.value for c in row] for row in ws_c.iter_rows(min_row=4, max_row=ws_c.max_row)]
        control = [r for r in circuit_rows if r[4] == "控制电缆出线"]
        self.assertTrue(control)
        self.assertEqual(len([r for r in control if r[1] == "02AP1"]), 6)  # 02AP1 的 6 条 KYJVP 回路
        # 控制电缆回路不写在元器件明细里，也不计入件数
        import re as _re
        mine = [r for r in rows if r[1] == "02AP1" and _re.fullmatch(r"\d+\.\d+", str(r[0] or ""))]
        self.assertEqual(sum(int(r[5] or 0) for r in mine), 6)


if __name__ == "__main__":
    unittest.main()
