# -*- coding: utf-8 -*-
"""组价明细单表的版面契约：导出 = 屏幕上那张组价大表。

锁住客户口径：只有一张「组价明细」表；顶部只有一行标题 + 一行费率口径；
列与组价页屏幕一致；行序分类行 → 箱柜行 → 明细行 → 构成行；
没有任何正式报价单抬头（封面/项目单位/报价人/审核人/报价说明/联系人）。
"""
import os
import tempfile
import unittest

import openpyxl

from extractor.excel import build_workbook, build_pricing_quotation_workbook
from extractor.quote_calc import quote_box, rmb_upper
from extractor.schema import Box, Circuit, Component, ExtractionResult

BANNED = ("封面", "报价人", "审核人", "报价说明", "项目单位", "联系人", "联系电话", "文件编号")


def _result() -> ExtractionResult:
    box = Box(code="02AL01", name="照明配电箱", install="挂墙明装", location="锅炉房",
              size="450*350*120", quantity=1, ip_rating="IP30")
    return ExtractionResult(
        title="测试项目",
        boxes=[box],
        circuits=[
            Circuit(box="02AL01", circuit_no="进线", breaker="SW-63A/3P", load_name="市电进线"),
            Circuit(box="02AL01", circuit_no="WL1", breaker="C9 2P C16A 6kA+ELE 30mA", load_name="照明"),
        ],
        components=[Component(name="空开", spec="C9 2P C16A 6kA+ELE 30mA", quantity=1, used_in="02AL01")],
        requirements=[],
        uncertainties=[],
    )


def _quote_res(quantity: int = 2) -> dict:
    """构造一份组价工作台核算结果：一行手填价、一行缺价、一行甲供不计价 + 壳体铜材。"""
    rows = [
        {"type": "box_header", "seq": "1", "box_code": "AP1", "name": "动力配电箱",
         "location": "车间", "design_spec": "GGD 800*2200*600",
         "replacement_model": "4种元器件明细", "quantity": quantity, "unit": "台",
         "unit_price": 776.86, "total_price": round(776.86 * quantity, 2), "source": "汇总"},
        {"type": "component", "seq": "1.1", "box_code": "AP1", "name": "微型断路器",
         "design_spec": "C16/1P", "replacement_model": "正泰 NXB-63 C16", "quantity": 3,
         "unit": "只", "unit_price": 12.34, "total_price": 37.02, "source": "手填",
         "brand": "正泰", "location": "主回路"},
        {"type": "component", "seq": "1.2", "box_code": "AP1", "name": "微型断路器（待确认）",
         "design_spec": "待确认", "replacement_model": "待指定报价型号", "quantity": 1,
         "unit": "只", "unit_price": 0.0, "total_price": 0.0, "source": "缺价",
         "brand": "待确认", "location": "主回路"},
        {"type": "component", "seq": "1.3", "box_code": "AP1", "name": "浪涌保护器",
         "design_spec": "SPD 40kA", "replacement_model": "甲供", "quantity": 1,
         "unit": "只", "unit_price": 0.0, "total_price": 0.0, "source": "甲供",
         "is_unpriced": True, "brand": "", "location": "主回路"},
        {"type": "material", "seq": "-", "box_code": "AP1", "name": "壳体",
         "design_spec": "GGD 800*2200*600", "replacement_model": "", "quantity": 1,
         "unit": "台", "unit_price": 500.0, "total_price": 500.0, "source": "定额",
         "brand": "成套定制", "location": ""},
        {"type": "material", "seq": "-", "box_code": "AP1", "name": "铜材",
         "design_spec": "母排 TMY 折算", "replacement_model": "", "quantity": 1,
         "unit": "套", "unit_price": 80.0, "total_price": 80.0, "source": "联动",
         "brand": "", "location": ""},
    ]
    breakdown = quote_box([
        {"unit_price": 12.34, "quantity": 3},
        {"unit_price": 0.0, "quantity": 1},   # 缺价行
        {"unit_price": 0.0, "quantity": 1, "priced": False},  # 甲供不计价
        {"unit_price": 500.0, "quantity": 1},
        {"unit_price": 80.0, "quantity": 1},
    ], quantity)
    return {
        "summary": {
            "brand": "正泰", "aux_rate": 0.05, "mfg_rate": 0.09, "tax_rate": 0.06,
            "total_price": breakdown["total"], "grand_total": int(round(breakdown["total"])),
            "total_upper": rmb_upper(int(round(breakdown["total"]))),
        },
        "rows": rows,
        "boxes": [{"box_code": "AP1", "name": "动力配电箱", "size": "GGD 800*2200*600",
                   "location": "车间", "quantity": quantity,
                   "unit_price": breakdown["unit_total"], "total_price": breakdown["total"],
                   "breakdown": breakdown}],
    }


class PricingDetailSheetTest(unittest.TestCase):
    def _export(self, result=None, **kwargs) -> openpyxl.Workbook:
        result = result or _result()
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
            path = tf.name
        try:
            build_workbook(result, "subtitle", path, layout="quote_template", **kwargs)
            return openpyxl.load_workbook(path)
        finally:
            if os.path.exists(path):
                os.remove(path)

    def _export_pricing(self, quote_res: dict | None = None) -> openpyxl.Workbook:
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
            path = tf.name
        try:
            build_pricing_quotation_workbook({"project": "某项目"}, quote_res or _quote_res(), path,
                                             target_brand="正泰")
            return openpyxl.load_workbook(path)
        finally:
            if os.path.exists(path):
                os.remove(path)

    def test_只有一张组价明细表(self):
        self.assertEqual(self._export().sheetnames[0], "组价明细")
        self.assertNotIn("组价汇总", self._export().sheetnames)
        self.assertEqual(self._export_pricing().sheetnames, ["组价明细"])

    def test_没有任何正式报价单抬头(self):
        for wb in (self._export(), self._export_pricing()):
            text = "\n".join(str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True)
                             for v in row if v)
            for banned in BANNED:
                self.assertNotIn(banned, text)

    def test_标题与表头行(self):
        ws = self._export_pricing().active
        self.assertEqual(ws["A1"].value, "某项目 组价明细")
        basis = str(ws["A2"].value)
        for need in ("正泰", "辅料率 5%", "成套制作费率 9%", "税费率 6%"):
            self.assertIn(need, basis)
        headers = [ws.cell(row=3, column=c).value for c in range(1, 13)]
        self.assertEqual(headers, ["序号", "柜号", "名称", "柜型 / 设计规格", "报价型号规格", "单位",
                                   "数量", "单价（元）", "总价（元）", "生产厂家", "价格来源", "备注"])
        self.assertEqual(ws.freeze_panes, "A4")

    def test_行序与屏幕一致(self):
        ws = self._export_pricing().active
        seqs = [ws.cell(row=r, column=1).value for r in range(4, ws.max_row + 1)]
        labels = [ws.cell(row=r, column=2).value for r in range(4, ws.max_row + 1)]
        # 分类行 → 箱柜行 → 明细行 → 构成行 → 合计 → 报价总计
        self.assertEqual(seqs[:2], ["一", "1-1"])
        self.assertEqual(seqs[2:5], ["1.1", "1.2", "1.3"])
        self.assertEqual(labels[7:12], ["小计", "辅料（小计 × 5%）", "成套制作费（(小计+辅料) × 9%）",
                                        "税费（前三项 × 6%）", "单台合计"])
        self.assertEqual(seqs[-2:], ["合计", "报价总计"])

    def test_缺价行单价总价为空_不计价行单价为0(self):
        ws = self._export_pricing().active
        rows = {ws.cell(row=r, column=1).value: r for r in range(4, ws.max_row + 1)}
        missing, unpriced = rows["1.2"], rows["1.3"]
        self.assertIsNone(ws.cell(row=missing, column=8).value)
        self.assertIsNone(ws.cell(row=missing, column=9).value)
        self.assertEqual(ws.cell(row=missing, column=11).value, "缺价")
        self.assertEqual(ws.cell(row=unpriced, column=8).value, 0.0)
        self.assertEqual(ws.cell(row=unpriced, column=9).value, 0.0)
        self.assertEqual(ws.cell(row=unpriced, column=11).value, "不计价")

    def test_构成行与合计取组价链结果(self):
        quote_res = _quote_res(quantity=2)
        breakdown = quote_res["boxes"][0]["breakdown"]
        ws = self._export_pricing(quote_res).active
        rows = {ws.cell(row=r, column=2).value: r for r in range(4, ws.max_row + 1)}
        self.assertAlmostEqual(float(ws.cell(row=rows["小计"], column=9).value), breakdown["subtotal"], places=2)
        self.assertAlmostEqual(float(ws.cell(row=rows["单台合计"], column=9).value), breakdown["unit_total"], places=2)
        total_row = [r for r in range(4, ws.max_row + 1) if ws.cell(row=r, column=1).value == "合计"][0]
        grand_row = total_row + 1
        self.assertAlmostEqual(float(ws.cell(row=total_row, column=9).value), breakdown["total"], places=2)
        self.assertEqual(int(ws.cell(row=grand_row, column=9).value), int(round(breakdown["total"])))
        self.assertIn(quote_res["summary"]["total_upper"], str(ws.cell(row=grand_row, column=3).value))


if __name__ == "__main__":
    unittest.main()
