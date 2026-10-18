# -*- coding: utf-8 -*-
"""组价导出的接口级契约：导出 = 组价页屏幕那张表，组价结果回写任务摘要。

覆盖 T15：
- 导出只有一张「组价明细」工作表，不含任何正式报价单抬头；
- 每台柜的五项构成、合计、报价总计与 GET /api/jobs/{id}/quotation 完全一致；
- 缺价行单价/总价为空并标“缺价”，不计价行单价 0；
- 真正请求过组价（GET/POST/导出）才把合计与总计写回 job.summary，
  拆量阶段导出不会把任务误标成已组价。
"""
import io
import unittest
import uuid

import openpyxl
import test_support  # noqa: F401
from fastapi.testclient import TestClient

from app import app as fastapi_app, jobs
from db import db_save_job

BANNED = ("封面", "报价人", "审核人", "报价说明", "项目单位", "联系人", "文件编号")


def _mk_job(**data_patch) -> str:
    job_id = "t15_" + uuid.uuid4().hex[:8]
    data = {
        "boxes": [
            {"code": "1AL1", "name": "照明配电箱", "type": "照明箱",
             "size": "450×350×120", "location": "锅炉房", "quantity": 2},
        ],
        "circuits": [
            {"box": "1AL1", "circuit_no": "进线", "breaker": "C65N-C40/3P"},
            {"box": "1AL1", "circuit_no": "WL1", "breaker": "C65N-C16/1P"},
        ],
        "components": [
            {"name": "微型断路器", "spec": "C65N-C16/1P", "unit": "只",
             "quantity": 4, "used_in": "1AL1"},
        ],
        "requirements": [],
        "uncertainties": [],
    }
    data.update(data_patch)
    jobs[job_id] = {
        "job_id": job_id,
        "filename": "组价导出测试.dwg",
        "file_type": "dwg",
        "status": "done",
        "project": "样例项目",
        "summary": {"title": "组价导出测试"},
        "data": data,
        "changes": [],
    }
    db_save_job(jobs[job_id])
    return job_id


class PricingExportTest(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(fastapi_app)

    def _export(self, job_id: str):
        resp = self.client.get(f"/api/jobs/{job_id}/excel", params={"stage": "pricing"})
        self.assertEqual(resp.status_code, 200)
        return openpyxl.load_workbook(io.BytesIO(resp.content))

    @staticmethod
    def _boxes(ws) -> dict:
        """按 柜号 切块：五项构成 + 该柜总价都取「总价」列（I列）。"""
        blocks, current = {}, None
        for r in range(1, ws.max_row + 1):
            code = ws.cell(row=r, column=2).value
            seq = ws.cell(row=r, column=1).value
            if isinstance(seq, str) and seq.startswith("1-") and code:
                current = {"code": code, "costs": {},
                           "quantity": ws.cell(row=r, column=7).value,
                           "total": ws.cell(row=r, column=9).value}
                blocks[code] = current
            elif current and isinstance(code, str) and code.split("（")[0] in (
                    "小计", "辅料", "成套制作费", "税费", "单台合计"):
                current["costs"][code.split("（")[0]] = float(ws.cell(row=r, column=9).value or 0)
        return blocks

    def test_导出只有一张组价明细且无正式报价单抬头(self):
        wb = self._export(_mk_job())
        self.assertEqual(wb.sheetnames, ["组价明细"])
        ws = wb.active
        self.assertTrue(str(ws["A1"].value).endswith("组价明细"))
        text = "\n".join(str(v) for row in ws.iter_rows(values_only=True) for v in row if v)
        for banned in BANNED:
            self.assertNotIn(banned, text)

    def test_每台柜构成与合计与屏幕接口完全一致(self):
        job_id = _mk_job()
        screen = self.client.get(f"/api/jobs/{job_id}/quotation").json()
        self.assertTrue(screen["ok"])
        ws = self._export(job_id).active
        blocks = self._boxes(ws)

        for box in screen["boxes"]:
            block = blocks[box["box_code"]]
            breakdown = box["breakdown"]
            self.assertEqual(int(block["quantity"]), int(box["quantity"]))
            self.assertAlmostEqual(block["costs"]["小计"], breakdown["subtotal"], places=2)
            self.assertAlmostEqual(block["costs"]["辅料"], breakdown["auxiliary"], places=2)
            self.assertAlmostEqual(block["costs"]["成套制作费"], breakdown["manufacture"], places=2)
            self.assertAlmostEqual(block["costs"]["税费"], breakdown["tax"], places=2)
            self.assertAlmostEqual(block["costs"]["单台合计"], breakdown["unit_total"], places=2)
            self.assertAlmostEqual(block["total"], breakdown["total"], places=2)

        rows = {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}
        self.assertAlmostEqual(float(ws.cell(row=rows["合计"], column=9).value),
                               screen["summary"]["total_price"], places=2)
        self.assertEqual(int(ws.cell(row=rows["报价总计"], column=9).value),
                         screen["summary"]["grand_total"])
        self.assertIn(screen["summary"]["total_upper"],
                      str(ws.cell(row=rows["报价总计"], column=3).value))

    def test_缺价行留空标缺价_不计价行单价为0(self):
        job_id = _mk_job(
            circuits=[{"box": "1AL1", "circuit_no": "进线", "breaker": "NM9-999/3P"}],
            components=[{"name": "浪涌保护器", "spec": "SPD-40kA", "unit": "只",
                         "quantity": 1, "used_in": "1AL1"}],
        )
        screen = self.client.get(f"/api/jobs/{job_id}/quotation").json()
        self.assertGreater(screen["summary"]["missing_count"], 0)
        ws = self._export(job_id).active
        seen_missing = seen_zero = False
        for r in range(4, ws.max_row + 1):
            if ws.cell(row=r, column=11).value == "缺价":
                self.assertIsNone(ws.cell(row=r, column=8).value)
                self.assertIsNone(ws.cell(row=r, column=9).value)
                seen_missing = True
            elif ws.cell(row=r, column=11).value == "不计价":
                self.assertEqual(ws.cell(row=r, column=8).value, 0.0)
                self.assertEqual(ws.cell(row=r, column=9).value, 0.0)
                seen_zero = True
        self.assertTrue(seen_missing, "缺价行应标“缺价”且单价/总价留空")
        self.assertTrue(seen_zero, "不计价行单价应为 0")

    def test_请求组价后合计与总计写回任务摘要(self):
        job_id = _mk_job()
        screen = self.client.get(f"/api/jobs/{job_id}/quotation").json()
        summary = jobs[job_id]["summary"]
        self.assertAlmostEqual(summary["total_price"], screen["summary"]["total_price"], places=2)
        self.assertEqual(summary["grand_total"], screen["summary"]["grand_total"])

        listed = {j["job_id"]: j for j in self.client.get("/api/jobs").json()["jobs"]}
        self.assertAlmostEqual(listed[job_id]["total_price"], screen["summary"]["total_price"], places=2)

    def test_保存手改价同样写回任务摘要(self):
        job_id = _mk_job()
        resp = self.client.post(f"/api/jobs/{job_id}/quotation",
                                json={"brand": "正泰", "overrides": {"C65N-C16/1P": {"unit_price": 99.9}}})
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertGreater(data["summary"]["total_price"], 0)
        self.assertEqual(jobs[job_id]["summary"]["total_price"], data["summary"]["total_price"])

    def test_组价导出也写回_拆量导出不误标已组价(self):
        job_id = _mk_job()
        takeoff = self.client.get(f"/api/jobs/{job_id}/excel", params={"stage": "takeoff"})
        self.assertEqual(takeoff.status_code, 200)
        self.assertIsNone(jobs[job_id]["summary"].get("total_price"),
                          "拆量阶段导出不应把任务标成已组价")

        self._export(job_id)
        self.assertGreater(jobs[job_id]["summary"]["total_price"], 0)
        self.assertGreater(jobs[job_id]["summary"]["grand_total"], 0)


if __name__ == "__main__":
    unittest.main()
