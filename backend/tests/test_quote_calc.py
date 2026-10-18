# -*- coding: utf-8 -*-
"""客户口径计费链：屏幕组价与导出 Excel 必须算出同一组数字。

客户报价单真值（制丝线升级改造项目）：
- 01LBZ1: 小计 11483.86 → 辅料 574.19 / 成套制作费 1085.22 / 税费 788.60 / 单台合计 13931.87
- 02AL01: 小计 2540.50 → 辅料 127.03 / 成套制作费 240.08 / 税费 174.46 / 单台合计 3082.07
- 项目合计 826288.83 → 总计 826289 → 捌拾贰万陆仟贰佰捌拾玖元整
"""
import os
import tempfile
import unittest
from decimal import Decimal, ROUND_HALF_UP

import openpyxl
from fastapi.testclient import TestClient

import test_support  # noqa: F401
from app import app, jobs
from db import db_save_job
from extractor.quote_calc import quote_box, quote_project, rmb_upper


class QuoteChainTest(unittest.TestCase):
    def _assert_box(self, lines, expected):
        quote = quote_box(lines, 1)
        for key, value in expected.items():
            self.assertAlmostEqual(quote[key], value, delta=0.01, msg=key)

    def test_01LBZ1_复算(self):
        self._assert_box([
            {"unit_price": 3200.0, "quantity": 2},
            {"unit_price": 2000.0, "quantity": 1},
            {"unit_price": 3083.86, "quantity": 1},
        ], {
            "subtotal": 11483.86,
            "auxiliary": 574.19,
            "manufacture": 1085.22,
            "tax": 788.60,
            "unit_total": 13931.87,
            "total": 13931.87,
        })

    def test_02AL01_复算(self):
        self._assert_box([
            {"unit_price": 1270.25, "quantity": 2},
        ], {
            "subtotal": 2540.5,
            "auxiliary": 127.03,
            "manufacture": 240.08,
            "tax": 174.46,
            "unit_total": 3082.07,
            "total": 3082.07,
        })

    def test_不计价行按0计入小计(self):
        quote = quote_box([
            {"unit_price": 1000, "quantity": 2},
            {"unit_price": 999, "quantity": 1, "priced": False},
        ])
        self.assertEqual(quote["subtotal"], 2000.0)
        self.assertEqual(quote["auxiliary"], 100.0)  # 2000 × 5%

    def test_柜总价等于单台合计乘台数(self):
        quote = quote_box([{"unit_price": 100, "quantity": 1}], quantity=3)
        self.assertAlmostEqual(quote["total"], quote["unit_total"] * 3, places=2)

    def test_项目合计总计与大写(self):
        project = quote_project([826288.83])
        self.assertAlmostEqual(project["total"], 826288.83, places=2)
        self.assertEqual(project["grand_total"], 826289)
        self.assertEqual(project["upper"], "捌拾贰万陆仟贰佰捌拾玖元整")
        self.assertEqual(int(Decimal("826288.83").quantize(Decimal("1"), rounding=ROUND_HALF_UP)),
                         project["grand_total"])

    def test_人民币大写(self):
        cases = {
            0: "零元整",
            10: "壹拾元整",
            1005: "壹仟零伍元整",
            100000: "壹拾万元整",
            826289: "捌拾贰万陆仟贰佰捌拾玖元整",
        }
        for amount, expected in cases.items():
            self.assertEqual(rmb_upper(amount), expected, amount)


class ScreenExportConsistencyTest(unittest.TestCase):
    """同一任务：GET /quotation 与 /excel?stage=pricing 的总计必须一致。"""

    def setUp(self):
        self.job_id = "test_quote_chain_001"
        self.job = {
            "id": self.job_id,
            "job_id": self.job_id,
            "filename": "计费链一致性测试.dwg",
            "file_type": "dwg",
            "status": "done",
            "progress": 100,
            "summary": {"title": "计费链一致性测试"},
            "data": {
                "boxes": [
                    {"code": "01LBZ1", "name": "照明配电总箱", "type": "GGD",
                     "size": "800×2200×600", "location": "制丝工房", "quantity": 1},
                    {"code": "02AL01", "name": "照明配电箱", "type": "照明箱",
                     "size": "450×350×120", "location": "锅炉房", "quantity": 2},
                ],
                "circuits": [
                    {"box": "01LBZ1", "circuit_no": "进线", "breaker": "NM1-250S/3300 200A"},
                    {"box": "02AL01", "circuit_no": "1WL1", "breaker": "C65N-C16/1P"},
                ],
                "components": [
                    {"name": "微型断路器", "spec": "C65N-C16/1P", "unit": "只",
                     "quantity": 4, "used_in": "01LBZ1"},
                    {"name": "电流互感器", "spec": "50/5", "unit": "只",
                     "quantity": 3, "used_in": "01LBZ1"},
                    {"name": "微型断路器", "spec": "C65N-C16/1P", "unit": "只",
                     "quantity": 6, "used_in": "02AL01"},
                ],
                "requirements": [],
                "uncertainties": [],
            },
        }
        jobs[self.job_id] = self.job
        db_save_job(self.job)
        self.client = TestClient(app)

    def _screen(self, **params):
        resp = self.client.get(f"/api/jobs/{self.job_id}/quotation", params=params)
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertTrue(data["ok"])
        return data

    def _detail_blocks(self, ws):
        """单表版式：1-1/1-2 箱柜行开块，小计/辅料/成套制作费/税费/单台合计 落在总价列。"""
        blocks = {}
        current = None
        for r in range(1, ws.max_row + 1):
            seq = ws.cell(row=r, column=1).value
            label = ws.cell(row=r, column=2).value
            if isinstance(seq, str) and seq.startswith("1-"):
                current = {"code": ws.cell(row=r, column=2).value, "costs": {}}
                blocks[current["code"]] = current
            elif current and isinstance(label, str) and label.split("（")[0] in (
                    "小计", "辅料", "成套制作费", "税费", "单台合计"):
                current["costs"][label.split("（")[0]] = float(ws.cell(row=r, column=9).value or 0)
        return blocks

    def test_屏幕与导出Excel总计一致(self):
        screen = self._screen()
        summary = screen["summary"]
        self.assertEqual(summary["grand_total"], int(Decimal(str(summary["total_price"]))
                                                      .quantize(Decimal("1"), rounding=ROUND_HALF_UP)))

        resp = self.client.get(f"/api/jobs/{self.job_id}/excel", params={"stage": "pricing"})
        self.assertEqual(resp.status_code, 200)
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
            path = tf.name
        try:
            with open(path, "wb") as f:
                f.write(resp.content)
            wb = openpyxl.load_workbook(path)
        finally:
            if os.path.exists(path):
                os.remove(path)

        # 组价导出只有一张表，与屏幕大表一一对应
        self.assertEqual(wb.sheetnames, ["组价明细"])
        ws = wb["组价明细"]
        rows = {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}
        total_row, grand_row = rows["合计"], rows["报价总计"]
        self.assertAlmostEqual(float(ws.cell(row=total_row, column=9).value),
                               summary["total_price"], places=2)
        self.assertEqual(int(ws.cell(row=grand_row, column=9).value), summary["grand_total"])
        self.assertIn(summary["total_upper"], str(ws.cell(row=grand_row, column=3).value))

        # 每个柜的五项费用也必须与屏幕同一结果
        blocks = self._detail_blocks(ws)
        for box in screen["boxes"]:
            detail = blocks[box["box_code"]]["costs"]
            breakdown = box["breakdown"]
            self.assertAlmostEqual(detail["小计"], breakdown["subtotal"], places=2)
            self.assertAlmostEqual(detail["辅料"], breakdown["auxiliary"], places=2)
            self.assertAlmostEqual(detail["成套制作费"], breakdown["manufacture"], places=2)
            self.assertAlmostEqual(detail["税费"], breakdown["tax"], places=2)
            self.assertAlmostEqual(detail["单台合计"], breakdown["unit_total"], places=2)

    def test_接口费率可覆盖并进入同一计费链(self):
        data = self._screen(aux_rate=0.05, mfg_rate=0.09, tax_rate=0.06)
        summary = data["summary"]
        self.assertEqual(summary["aux_rate"], 0.05)
        self.assertEqual(summary["mfg_rate"], 0.09)
        self.assertEqual(summary["tax_rate"], 0.06)
        for box in data["boxes"]:
            breakdown = box["breakdown"]
            self.assertEqual(breakdown["rates"], {"aux": 0.05, "mfg": 0.09, "tax": 0.06})
            self.assertAlmostEqual(
                breakdown["subtotal"] + breakdown["auxiliary"]
                + breakdown["manufacture"] + breakdown["tax"],
                breakdown["unit_total"], places=2)

        changed = self._screen(aux_rate=0.08, mfg_rate=0.12, tax_rate=0.10)
        self.assertEqual(changed["summary"]["aux_rate"], 0.08)
        self.assertGreater(changed["summary"]["total_price"], summary["total_price"])


if __name__ == "__main__":
    unittest.main()
