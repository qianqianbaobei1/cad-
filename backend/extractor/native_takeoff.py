# -*- coding: utf-8 -*-
"""图纸原生读取与数据导出（不经视觉模型）。

参照《CAD快速看图》的读取能力补齐我们的原生链路，见
docs/逆向分析_CAD快速看图_兼容性与读取能力.md。它对每类对象单独解析，
所以框选、搜索、表格提取不丢内容；我们这里按同样的思路，把图纸里能直接读到的
文字、表格、图块数量、线长、字体情况一次读全，并能导出成 Excel。

读取优先级（高 → 低，前一级读不到才用后一级）：

DWG / DXF
  1. 原生对象：TEXT / MTEXT / 块内文字 / 块属性 / MINSERT 阵列 / 尺寸标注显示文字 /
     多重引线 / 表格（ACAD_TABLE 的图形块）/ 代理对象的代理图形（天正、浩辰等自定义对象）
  2. DWG 转换由 cad.dwg_to_dxf 负责（LibreDWG → acadrust），转换报错在概览里如实列出
  3. 原生内容不足时，由 app.py 的准入门禁回退到渲染 + 视觉模型
PDF
  1. PDF 文字层（CAD 直接打印的 PDF 通常带）
  2. 文字层为空（文字转曲）→ 本机 OCR（tesseract，可选）
  3. 仍不足 → 视觉模型（app.py 现有流程）

“量”的原生来源：
  - 图块统计：同名块引用次数 = 图例/设备数量（MINSERT 按行×列计），不依赖块是否带属性
  - 图层线长：按图层累计线、多段线、圆弧、样条长度（电缆、桥架、管线）
  - 表格：系统图里的 ACAD_TABLE 与线框表格，按网格还原成行列
"""
from __future__ import annotations

import json
import math
import os
import re
import shutil
import subprocess
import tempfile
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from typing import Any

from .cad import clean_mtext, load_dxf_document, resolve_dwg_dxf, sanitize_surrogates
from .tangent import safe_decode_tangent_if_cipher

MAX_DEPTH = 6
MAX_TEXTS = 300_000

# 天正对象常见的块名前缀（天正电气 / 弱电 / 设备），只用于在统计里打标签
_TANGENT_BLOCK = re.compile(r"^\$(EQUIP|TelecSys|element|equip|TCH|TEL)\$", re.I)
# 天正自定义对象的类名（TDbText / TDbEquipment / TCH_TEXT …），其文字可能是 26 进制加密串
_TANGENT_CLASS = re.compile(r"^(TDb|TCH_|TEL_)", re.I)

_SPACE_MODEL = "模型空间"


@dataclass
class TextRecord:
    text: str
    x: float
    y: float
    height: float
    layer: str
    kind: str              # TEXT / MTEXT / ATTRIB / DIMENSION / MLEADER / TABLE / PROXY / PDF_TEXT / OCR
    space: str = _SPACE_MODEL
    owner: str = ""        # 所在块名 / 表格编号 / 代理类名
    handle: str = ""
    style: str = ""
    table: int | None = None


@dataclass
class NativeInventory:
    source: str
    file_type: str
    texts: list[TextRecord] = field(default_factory=list)
    tables: list[dict[str, Any]] = field(default_factory=list)
    blocks: list[dict[str, Any]] = field(default_factory=list)
    layer_lengths: list[dict[str, Any]] = field(default_factory=list)
    fonts: list[dict[str, Any]] = field(default_factory=list)
    coverage: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    read_route: list[str] = field(default_factory=list)
    units: str = ""

    def summary(self) -> dict[str, Any]:
        kinds = Counter(t.kind for t in self.texts)
        return {
            "source": os.path.basename(self.source),
            "file_type": self.file_type,
            "read_route": self.read_route,
            "units": self.units,
            "texts": len(self.texts),
            "texts_by_kind": dict(kinds),
            "tables": len(self.tables),
            "block_kinds": len(self.blocks),
            "block_instances": sum(b["count"] for b in self.blocks),
            "top_blocks": [{k: b[k] for k in ("name", "count", "sample_text", "tangent")}
                           for b in self.blocks[:15]],
            "layers_with_length": len(self.layer_lengths),
            "fonts_missing": [f for f in self.fonts if not f["font_found"] or
                              (f["bigfont"] and not f["bigfont_found"])],
            "coverage": self.coverage,
            "warnings": self.warnings,
        }

    def to_json(self) -> dict[str, Any]:
        data = asdict(self)
        data["summary"] = self.summary()
        return data


# ---------------------------------------------------------------- 文字几何

def _text_width(text: str, height: float, width_factor: float = 1.0) -> float:
    """估算单行文字宽度：中文按 1 个字高，西文按 0.6 个字高。"""
    w = sum(1.0 if ord(ch) > 0x2E7F else 0.6 for ch in text)
    return w * height * (width_factor or 1.0)


def _text_center(e: Any, kind: str, text: str) -> tuple[float, float, float]:
    """文字的几何中心（局部坐标）与字高。表格落格、框选按中心判断，不能用插入基点。"""
    if kind == "MTEXT":
        h = float(e.dxf.get("char_height", 2.5) or 2.5)
        lines = text.split("\n")
        w = float(e.dxf.get("width", 0) or 0) or max(_text_width(s, h) for s in lines)
        total_h = h * (1.0 + 1.4 * (len(lines) - 1))
        ap = int(e.dxf.get("attachment_point", 1) or 1)
        ix, iy = float(e.dxf.insert.x), float(e.dxf.insert.y)
        col, row = (ap - 1) % 3, (ap - 1) // 3          # col: 0 左 1 中 2 右；row: 0 上 1 中 2 下
        cx = ix + (0.5 - 0.5 * col) * w
        cy = iy + (-0.5 + 0.5 * row) * total_h
        return cx, cy, h
    h = float(e.dxf.get("height", 2.5) or 2.5)
    w = _text_width(text, h, float(e.dxf.get("width", 1.0) or 1.0))
    halign = int(e.dxf.get("halign", 0) or 0)
    valign = int(e.dxf.get("valign", 0) or 0)
    if (halign or valign) and e.dxf.hasattr("align_point"):
        ax, ay = float(e.dxf.align_point.x), float(e.dxf.align_point.y)
    else:
        ax, ay = float(e.dxf.insert.x), float(e.dxf.insert.y)
    if halign in (1, 3, 4, 5):
        cx = ax
    elif halign == 2:
        cx = ax - w / 2
    else:
        cx = ax + w / 2
    if halign == 4 or valign == 2:
        cy = ay
    elif valign == 3:
        cy = ay - h / 2
    else:
        cy = ay + h / 2
    rot = math.radians(float(e.dxf.get("rotation", 0) or 0))
    if rot and kind != "MTEXT":
        dx, dy = cx - ax, cy - ay
        cx = ax + dx * math.cos(rot) - dy * math.sin(rot)
        cy = ay + dx * math.sin(rot) + dy * math.cos(rot)
    return cx, cy, h


def _entity_text(e: Any, kind: str) -> str:
    try:
        if kind == "MTEXT":
            raw = e.plain_text()
        else:
            raw = e.dxf.get("text", "")
    except Exception:
        raw = getattr(e, "text", "") or ""
    return sanitize_surrogates(clean_mtext(str(raw or "")))


# ---------------------------------------------------------------- DWG / DXF 遍历

class _Walker:
    """把一张图里所有带文字的对象读成 TextRecord（模型空间坐标）。

    块定义按块名缓存局部坐标结果，引用时只做一次矩阵变换——快速看图同样是先解析块再实例化，
    1 万多个块引用也只展开一次块定义。
    """

    def __init__(self, doc: Any):
        self.doc = doc
        self.cache: dict[str, list[dict[str, Any]]] = {}
        self.proxy_without_graphics = 0
        self.proxy_classes: Counter = Counter()
        self.tangent_decrypted = 0

    def _tangent(self, items: list[dict[str, Any]]) -> list[dict[str, Any]]:
        """天正对象 / 天正块里的文字若是加密串就解出明文（tangent.py，未知版本按默认分支）。"""
        out = []
        for it in items:
            plain = safe_decode_tangent_if_cipher(it["text"])
            if plain != it["text"]:
                self.tangent_decrypted += 1
                it = {**it, "text": plain, "cipher": it["text"]}
            out.append(it)
        return out

    # 局部坐标下的文字（块定义内部，含嵌套块、标注、代理图形）
    def local(self, name: str, depth: int = 0) -> list[dict[str, Any]]:
        if name in self.cache:
            return self.cache[name]
        self.cache[name] = []
        blk = self.doc.blocks.get(name)
        out: list[dict[str, Any]] = []
        if blk is not None and depth <= MAX_DEPTH:
            for e in blk:
                out.extend(self._entity(e, depth + 1, owner=name))
        if _TANGENT_BLOCK.match(name):
            out = self._tangent(out)
        self.cache[name] = out
        return out

    def _rec(self, e: Any, kind: str, text: str, owner: str, **extra) -> dict[str, Any]:
        cx, cy, h = _text_center(e, "MTEXT" if kind in ("MTEXT", "MLEADER") else "TEXT", text) \
            if kind in ("TEXT", "MTEXT", "ATTRIB") else (extra.pop("x"), extra.pop("y"), extra.pop("h", 2.5))
        return {"text": text, "x": cx, "y": cy, "height": h,
                "layer": str(e.dxf.get("layer", "0")), "kind": kind, "owner": owner,
                "handle": str(e.dxf.get("handle", "")), "style": str(e.dxf.get("style", "")) if e.dxf.is_supported("style") else "", **extra}

    def _transformed(self, items: list[dict[str, Any]], m: Any, owner: str) -> list[dict[str, Any]]:
        # 字高跟着插入比例缩放：比例 100 的图框块里 3.5 高的字，图面上是 350
        scale = float(m.transform_direction((0.0, 1.0, 0.0)).magnitude) or 1.0
        res = []
        for it in items:
            p = m.transform((it["x"], it["y"], 0.0))
            res.append({**it, "x": float(p[0]), "y": float(p[1]), "owner": it["owner"] or owner,
                        "height": float(it.get("height") or 0) * scale})
        return res

    def _entity(self, e: Any, depth: int, owner: str = "") -> list[dict[str, Any]]:
        kind = e.dxftype()
        try:
            if kind in ("TEXT", "MTEXT"):
                text = _entity_text(e, kind)
                return [self._rec(e, kind, text, owner)] if text else []
            if kind in ("INSERT",):
                return self._insert(e, depth, owner)
            if kind == "DIMENSION":
                return self._dimension(e, depth, owner)
            if kind in ("MULTILEADER", "MLEADER"):
                return self._mleader(e, owner)
            if kind == "ACAD_TABLE":
                return self._table_texts(e, depth)
            if kind == "ACAD_PROXY_ENTITY":
                return self._proxy(e, depth, owner)
        except Exception:
            return []
        return []

    def _insert(self, e: Any, depth: int, owner: str) -> list[dict[str, Any]]:
        out: list[dict[str, Any]] = []
        name = str(e.dxf.name)
        rows = int(e.dxf.get("row_count", 1) or 1)
        cols = int(e.dxf.get("column_count", 1) or 1)
        targets = list(e.multi_insert()) if rows * cols > 1 else [e]
        local = self.local(name, depth) if depth <= MAX_DEPTH else []
        for ins in targets:
            try:
                m = ins.matrix44()
            except Exception:
                continue
            out.extend(self._transformed(local, m, owner or name))
        attribs = []
        for a in getattr(e, "attribs", []):
            text = _entity_text(a, "TEXT")
            if text:
                attribs.append(self._rec(a, "ATTRIB", text, owner or name, tag=str(a.dxf.get("tag", ""))))
        out.extend(self._tangent(attribs) if _TANGENT_BLOCK.match(name) else attribs)
        return out

    def _dimension(self, e: Any, depth: int, owner: str) -> list[dict[str, Any]]:
        blk = None
        try:
            blk = e.get_geometry_block()
        except Exception:
            blk = None
        if blk is not None:
            items = [it for it in self.local(blk.name, depth) if it["kind"] in ("TEXT", "MTEXT")]
            if items:
                return [{**it, "kind": "DIMENSION", "owner": owner, "handle": str(e.dxf.get("handle", ""))}
                        for it in items]
        text = str(e.dxf.get("text", "") or "")
        if text and text != "<>":
            p = e.dxf.get("text_midpoint") or e.dxf.get("defpoint")
            if p is not None:
                return [self._rec(e, "DIMENSION", clean_mtext(text).replace("<>", ""), owner,
                                  x=float(p[0]), y=float(p[1]), h=2.5)]
        return []

    def _mleader(self, e: Any, owner: str) -> list[dict[str, Any]]:
        ctx = getattr(e, "context", None)
        mt = getattr(ctx, "mtext", None) if ctx is not None else None
        if mt is None:
            return []
        text = sanitize_surrogates(clean_mtext(str(getattr(mt, "default_content", "") or "")))
        if not text:
            return []
        ins = getattr(mt, "insert", None)
        h = float(getattr(ctx, "char_height", 2.5) or 2.5)
        if ins is None:
            return []
        return [self._rec(e, "MLEADER", text, owner, x=float(ins[0]), y=float(ins[1]), h=h)]

    def _table_texts(self, e: Any, depth: int) -> list[dict[str, Any]]:
        name = e.get_block_name()
        if not name:
            return []
        loc = e.get_insert_location()
        ox, oy = float(loc[0]), float(loc[1])
        return [{**it, "x": it["x"] + ox, "y": it["y"] + oy, "kind": "TABLE", "owner": name}
                for it in self.local(name, depth)]

    def _proxy(self, e: Any, depth: int, owner: str) -> list[dict[str, Any]]:
        cls = ""
        try:
            cls = str(e.dxf.get("original_class_name", "") or e.dxf.get("proxy_entity_class_id", ""))
        except Exception:
            pass
        self.proxy_classes[cls or "未知类"] += 1
        try:
            virtual = list(e.virtual_entities())
        except Exception:
            virtual = []
        if not virtual:
            self.proxy_without_graphics += 1
            return []
        out = []
        for v in virtual:
            for it in self._entity(v, depth + 1, owner):
                out.append({**it, "kind": "PROXY" if it["kind"] in ("TEXT", "MTEXT") else it["kind"],
                            "owner": cls or owner})
        return self._tangent(out) if _TANGENT_CLASS.match(cls) else out

    def layout_texts(self, layout: Any, space: str) -> list[TextRecord]:
        res: list[TextRecord] = []
        for e in layout:
            for it in self._entity(e, 0):
                res.append(TextRecord(text=it["text"], x=round(it["x"], 3), y=round(it["y"], 3),
                                      height=round(float(it.get("height") or 0), 3),
                                      layer=sanitize_surrogates(it.get("layer", "")), kind=it["kind"],
                                      space=space, owner=sanitize_surrogates(it.get("owner", "")),
                                      handle=it.get("handle", ""), style=it.get("style", "")))
                if len(res) >= MAX_TEXTS:
                    return res
        return res


def _effective_block_name(doc: Any, name: str) -> str:
    """动态块的匿名块（*U…）找回原始块名，否则同一种设备会被拆成很多行。"""
    if not name.startswith("*U"):
        return name
    try:
        br = doc.block_records.get(name)
        if br is not None and br.has_xdata("AcDbBlockRepBTag"):
            for code, value in br.get_xdata("AcDbBlockRepBTag"):
                if code == 1005:
                    orig = doc.entitydb.get(value)
                    if orig is not None:
                        return str(orig.dxf.name)
    except Exception:
        pass
    return name


def block_statistics(doc: Any, walker: _Walker) -> list[dict[str, Any]]:
    """模型空间顶层块引用统计：同名块数量 = 图例/设备数量（MINSERT 按行×列计）。"""
    from ezdxf import bbox as _bbox

    stats: dict[str, dict[str, Any]] = {}
    for e in doc.modelspace().query("INSERT"):
        raw = str(e.dxf.name)
        name = _effective_block_name(doc, raw)
        n = int(e.dxf.get("row_count", 1) or 1) * int(e.dxf.get("column_count", 1) or 1)
        s = stats.setdefault(name, {"name": name, "raw_names": set(), "count": 0,
                                    "layers": Counter(), "attribs": defaultdict(Counter)})
        s["raw_names"].add(raw)
        s["count"] += n
        s["layers"][str(e.dxf.get("layer", "0"))] += n
        for a in getattr(e, "attribs", []):
            v = clean_mtext(str(a.dxf.get("text", "") or ""))
            if v and _TANGENT_BLOCK.match(name):
                v = safe_decode_tangent_if_cipher(v)
            if v:
                s["attribs"][str(a.dxf.get("tag", ""))][v] += n
    rows = []
    for name, s in stats.items():
        sample_name = sorted(s["raw_names"])[0]
        inner = []
        for it in walker.local(sample_name):
            if it["text"] not in inner:
                inner.append(it["text"])
            if len(inner) >= 6:
                break
        size = ""
        try:
            ext = _bbox.extents(doc.blocks.get(sample_name), fast=True)
            if ext.has_data:
                size = f"{ext.size.x:.0f}×{ext.size.y:.0f}"
        except Exception:
            pass
        attrs = "；".join(f"{tag}={'/'.join(v for v, _ in c.most_common(3))}"
                         for tag, c in list(s["attribs"].items())[:4])
        rows.append({
            "name": sanitize_surrogates(name),
            "anonymous": name.startswith("*"),
            "tangent": bool(_TANGENT_BLOCK.match(name)),
            "count": s["count"],
            "layers": "、".join(sanitize_surrogates(k) for k, _ in s["layers"].most_common(3)),
            "sample_text": sanitize_surrogates(" / ".join(inner)),
            "attributes": sanitize_surrogates(attrs),
            "size": size,
        })
    rows.sort(key=lambda r: (-r["count"], r["name"]))
    return rows


def _curve_length(e: Any) -> float:
    kind = e.dxftype()
    if kind == "LINE":
        return float((e.dxf.end - e.dxf.start).magnitude)
    if kind == "ARC":
        span = (float(e.dxf.end_angle) - float(e.dxf.start_angle)) % 360.0
        return math.radians(span) * float(e.dxf.radius)
    if kind == "CIRCLE":
        return 2 * math.pi * float(e.dxf.radius)
    if kind in ("LWPOLYLINE", "POLYLINE", "SPLINE", "ELLIPSE"):
        from ezdxf import path as _path
        pts = list(_path.make_path(e).flattening(distance=0.5))
        return sum((pts[i + 1] - pts[i]).magnitude for i in range(len(pts) - 1))
    return 0.0


def layer_lengths(doc: Any) -> list[dict[str, Any]]:
    """按图层累计模型空间顶层线长（不进块内部，避免把图例符号的线算进电缆长度）。"""
    acc: dict[str, dict[str, Any]] = {}
    for e in doc.modelspace().query("LINE ARC CIRCLE LWPOLYLINE POLYLINE SPLINE ELLIPSE"):
        try:
            length = _curve_length(e)
        except Exception:
            continue
        if length <= 0:
            continue
        layer = sanitize_surrogates(str(e.dxf.get("layer", "0")))
        a = acc.setdefault(layer, {"layer": layer, "count": 0, "length": 0.0, "types": Counter()})
        a["count"] += 1
        a["length"] += length
        a["types"][e.dxftype()] += 1
    rows = [{"layer": a["layer"], "count": a["count"], "length": round(a["length"], 2),
             "types": "、".join(f"{k}×{v}" for k, v in a["types"].most_common())}
            for a in acc.values()]
    rows.sort(key=lambda r: -r["length"])
    return rows


# ---------------------------------------------------------------- 字体

def _font_dirs() -> list[str]:
    dirs = []
    env = os.environ.get("EXTRACTOR_FONT_DIRS", "")
    dirs += [d for d in env.split(os.pathsep) if d]
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    dirs += [os.path.join(here, "fonts"), os.path.expanduser("~/Library/Fonts"), "/Library/Fonts",
             "/System/Library/Fonts", "/System/Library/Fonts/Supplemental",
             "/usr/share/fonts", "C:\\Windows\\Fonts"]
    try:
        import ezdxf
        dirs += list(ezdxf.options.support_dirs)
    except Exception:
        pass
    return [d for d in dirs if os.path.isdir(d)]


_FONT_INDEX: dict[str, str] | None = None


def _font_index() -> dict[str, str]:
    global _FONT_INDEX
    if _FONT_INDEX is None:
        idx: dict[str, str] = {}
        for d in _font_dirs():
            for root, _, files in os.walk(d):
                for f in files:
                    idx.setdefault(f.lower(), os.path.join(root, f))
        _FONT_INDEX = idx
    return _FONT_INDEX


def _find_font(name: str) -> str:
    name = (name or "").strip()
    if not name:
        return ""
    idx = _font_index()
    low = os.path.basename(name).lower()
    cands = [low] if os.path.splitext(low)[1] else [low + ".shx", low + ".ttf", low + ".ttc"]
    for c in cands:
        if c in idx:
            return idx[c]
    return ""


def font_report(doc: Any, texts: list[TextRecord]) -> list[dict[str, Any]]:
    """每个文字样式用的字体 / 大字体能否在本机找到。找不到的会被替换，字形就与原图不同。"""
    usage = Counter(t.style.lower() for t in texts if t.style)
    rows = []
    for s in doc.styles:
        name = str(s.dxf.name)
        font = str(s.dxf.get("font", "") or "")
        big = str(s.dxf.get("bigfont", "") or "")
        font_path = _find_font(font)
        big_path = _find_font(big) if big else ""
        rows.append({
            "style": sanitize_surrogates(name), "font": font, "bigfont": big,
            "font_found": bool(font_path) or not font, "bigfont_found": bool(big_path),
            "font_path": font_path, "bigfont_path": big_path,
            "used": usage.get(name.lower(), 0),
        })
    rows.sort(key=lambda r: -r["used"])
    return rows


# ---------------------------------------------------------------- 表格

def _grid_from(lines_h: list, lines_v: list, texts: list[dict[str, Any]], min_h: float):
    from .gtable import build_relation_grid, load_texts_into_cells, merge_lines
    h = merge_lines(lines_h, True, min_h)
    v = merge_lines(lines_v, False, min_h)
    grid = build_relation_grid(h, v, min_text_height=min_h)
    if not grid:
        return None
    # load_texts_into_cells 以 (x + w/2, y + h/2) 作中心；这里传入的是中心点，换算回去
    shifted = [{"text": t["text"], "x": t["x"] - t["w"] / 2, "y": t["y"] - t["h"] / 2,
                "width": t["w"], "height": t["h"]} for t in texts]
    load_texts_into_cells(grid, shifted, tol=min_h * 0.1)
    return grid


def _compact(matrix: list[list[str]]) -> list[list[str]]:
    """去掉全空的行和列（线框里常混进图框线、回路线，切出大量空列）。"""
    rows = [r for r in matrix if any(c.strip() for c in r)]
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    keep = [j for j in range(width) if any(r[j].strip() for r in rows)]
    return [[r[j] for j in keep] for r in rows]


def _fill_ratio(matrix: list[list[str]]) -> float:
    cells = sum(len(r) for r in matrix)
    return sum(1 for r in matrix for c in r if c.strip()) / cells if cells else 0.0


def _rows_by_y(texts: list[dict[str, Any]], tol: float) -> list[list[str]]:
    """没有网格线时的兜底：按 y 聚成行，行内按 x 排序。"""
    rows: list[list[dict[str, Any]]] = []
    for t in sorted(texts, key=lambda t: -t["y"]):
        if rows and abs(rows[-1][0]["y"] - t["y"]) <= tol:
            rows[-1].append(t)
        else:
            rows.append([t])
    return [[t["text"] for t in sorted(r, key=lambda t: t["x"])] for r in rows]


def table_from_block(doc: Any, walker: _Walker, block_name: str) -> list[list[str]]:
    """ACAD_TABLE 的图形块（*T…）还原成行列。表格对象本身被转换器读坏时，图形块往往仍完好。"""
    from .gtable import extract_hv_lines
    blk = doc.blocks.get(block_name)
    if blk is None:
        return []
    texts = [{"text": it["text"], "x": it["x"], "y": it["y"], "h": float(it["height"] or 2.5),
              "w": _text_width(it["text"], float(it["height"] or 2.5))}
             for it in walker.local(block_name)]
    if not texts:
        return []
    min_h = max(min(t["h"] for t in texts), 0.5)
    lines = [e for e in blk if e.dxftype() in ("LINE", "LWPOLYLINE", "POLYLINE")]
    try:
        h, v = extract_hv_lines(lines)
        grid = _grid_from(h, v, texts, min_h) if len(h) >= 2 and len(v) >= 2 else None
    except Exception:
        grid = None
    if grid:
        matrix = _compact([[c.text_content for c in row] for row in grid])
        if sum(1 for row in matrix for c in row if c) >= max(1, len(texts) // 2):
            return matrix
    return _rows_by_y(texts, min_h * 0.6)


def collect_tables(doc: Any, walker: _Walker) -> list[dict[str, Any]]:
    tables = []
    for layout in doc.layouts:
        for e in layout.query("ACAD_TABLE"):
            try:
                name = e.get_block_name()
                loc = e.get_insert_location()
                matrix = table_from_block(doc, walker, name)
            except Exception:
                continue
            if not matrix:
                continue
            tables.append({
                "index": len(tables) + 1, "source": "ACAD_TABLE", "block": name,
                "space": _SPACE_MODEL if layout.name == "Model" else layout.name,
                "x": round(float(loc[0]), 2), "y": round(float(loc[1]), 2),
                "rows": len(matrix), "cols": max(len(r) for r in matrix), "cells": matrix,
            })
    return tables


# ---------------------------------------------------------------- 覆盖统计

_HANDLED = {
    "TEXT": "文字", "MTEXT": "文字", "INSERT": "块（含属性、块内文字、MINSERT 阵列）",
    "DIMENSION": "标注显示文字", "MULTILEADER": "多重引线文字", "ACAD_TABLE": "表格（从图形块还原）",
    "ACAD_PROXY_ENTITY": "代理对象（炸开代理图形）",
    "LINE": "几何 / 线长", "LWPOLYLINE": "几何 / 线长", "POLYLINE": "几何 / 线长", "ARC": "几何 / 线长",
    "CIRCLE": "几何 / 线长", "SPLINE": "几何 / 线长", "ELLIPSE": "几何 / 线长",
    "HATCH": "几何", "SOLID": "几何", "POINT": "几何", "3DFACE": "几何", "ATTDEF": "块定义模板（不计）",
    "VIEWPORT": "视口", "WIPEOUT": "遮罩",
}


def entity_coverage(doc: Any) -> list[dict[str, Any]]:
    cnt: Counter = Counter()
    for layout in doc.layouts:
        for e in layout:
            cnt[e.dxftype()] += 1
    return [{"type": k, "count": v, "handling": _HANDLED.get(k, "未解析（仅可能在渲染中出现）")}
            for k, v in cnt.most_common()]


def _conversion_notes(dxf_path: str) -> list[str]:
    notes = []
    path = dxf_path + ".conversion.json"
    if not os.path.isfile(path):
        return notes
    try:
        rec = json.load(open(path, encoding="utf-8"))
    except Exception:
        return notes
    notes.append(f"DWG 转换：{rec.get('converter') or rec.get('tool')}，"
                 f"警告 {rec.get('warning_lines', 0)} 行、错误 {rec.get('error_lines', 0)} 行"
                 + ("（已用备用转换器）" if rec.get("fallback_used") else ""))
    excerpt = str(rec.get("stderr_excerpt") or "")
    if "ACAD_TABLE" in excerpt:
        notes.append("转换器读表格对象（ACAD_TABLE）报错，表格内容已改从表格图形块读取")
    return notes


# ---------------------------------------------------------------- 入口

def scan_cad(path: str) -> NativeInventory:
    """读取 DWG / DXF 的全部原生内容。"""
    is_dwg = path.lower().endswith(".dwg")
    dxf_path = resolve_dwg_dxf(path) if is_dwg else path
    doc = load_dxf_document(dxf_path)
    inv = NativeInventory(source=path, file_type="dwg" if is_dwg else "dxf")
    inv.read_route.append("DWG → DXF 转换后原生读取" if is_dwg else "DXF 原生读取")
    if is_dwg:
        inv.warnings.extend(_conversion_notes(dxf_path))
    units = {0: "未设置（按毫米折算）", 1: "英寸", 4: "毫米", 5: "厘米", 6: "米"}
    inv.units = units.get(int(doc.header.get("$INSUNITS", 0) or 0), str(doc.header.get("$INSUNITS")))

    walker = _Walker(doc)
    inv.texts = walker.layout_texts(doc.modelspace(), _SPACE_MODEL)
    for layout in doc.layouts:
        if layout.name != "Model":
            inv.texts.extend(walker.layout_texts(layout, f"布局：{layout.name}"))
    inv.tables = collect_tables(doc, walker)
    inv.blocks = block_statistics(doc, walker)
    inv.layer_lengths = layer_lengths(doc)
    inv.fonts = font_report(doc, inv.texts)
    inv.coverage = entity_coverage(doc)

    if walker.proxy_classes:
        inv.warnings.append("代理对象（自定义对象）"
                            + "、".join(f"{k}×{v}" for k, v in walker.proxy_classes.most_common(8)))
    if walker.tangent_decrypted:
        inv.warnings.append(f"已解密 {walker.tangent_decrypted} 条天正加密文字（按默认密钥分支，请抽查）")
    if walker.proxy_without_graphics:
        inv.warnings.append(f"{walker.proxy_without_graphics} 个代理对象没有保存代理图形，"
                            "其内容无法原生读取（原软件存图时关闭了代理图形）")
    missing = [f for f in inv.fonts if f["used"] and (not f["font_found"] or (f["bigfont"] and not f["bigfont_found"]))]
    if missing:
        inv.warnings.append("本机缺少字体（显示会被替换）：" + "、".join(
            sorted({x for f in missing for x in (f["font"], f["bigfont"]) if x and not _find_font(x)})))
    return inv


def _tesseract() -> str | None:
    return shutil.which("tesseract")


OCR_LONG_EDGE_PX = 9000     # 整页渲染到这个长边；A1 缩印成 A4 的图，1mm 高的字约 30px
OCR_TILE_PX = 2600          # 分块识别，块间重叠 12%
OCR_MIN_CONF = 55
_MEANINGFUL = re.compile(r"[A-Za-z0-9\u4e00-\u9fff]")


def _ocr_tile(exe: str, img_path: str) -> str:
    try:
        return subprocess.run([exe, img_path, "-", "-l", "chi_sim+eng", "--psm", "11", "tsv"],
                              capture_output=True, text=True, timeout=600).stdout
    except (subprocess.TimeoutExpired, OSError):
        return ""


def _ocr_page(page: Any) -> list[TextRecord]:
    """文字转曲的 PDF 用本机 tesseract 识别（只做兜底，结果标为 OCR，需人工核对）。

    小字靠放大：整页按长边约 9000px 渲染再分块识别，和视觉模型“切块放大”同一个道理。
    """
    exe = _tesseract()
    if not exe:
        return []
    from concurrent.futures import ThreadPoolExecutor
    long_pt = max(float(page.rect.width), float(page.rect.height))
    dpi = int(min(900, max(300, OCR_LONG_EDGE_PX * 72.0 / long_pt)))
    scale = 72.0 / dpi
    ph = float(page.rect.height)
    out: list[TextRecord] = []
    seen: set = set()
    with tempfile.TemporaryDirectory() as tmp:
        pix = page.get_pixmap(dpi=dpi)
        W, H = pix.width, pix.height
        step = int(OCR_TILE_PX * 0.88)
        tiles = []
        import pymupdf
        for y0 in range(0, max(1, H - int(OCR_TILE_PX * 0.12)), step):
            for x0 in range(0, max(1, W - int(OCR_TILE_PX * 0.12)), step):
                x1, y1 = min(W, x0 + OCR_TILE_PX), min(H, y0 + OCR_TILE_PX)
                clip = pymupdf.Rect(x0 * scale, y0 * scale, x1 * scale, y1 * scale)
                path = os.path.join(tmp, f"t{x0}_{y0}.png")
                page.get_pixmap(dpi=dpi, clip=clip).save(path)
                tiles.append((x0, y0, path))
        with ThreadPoolExecutor(max_workers=4) as pool:
            results = list(pool.map(lambda t: (t[0], t[1], _ocr_tile(exe, t[2])), tiles))
    for x0, y0, tsv in results:
        for line in tsv.splitlines()[1:]:
            parts = line.split("\t")
            if len(parts) < 12 or not parts[11].strip():
                continue
            try:
                conf = float(parts[10])
                left, top, w, h = (int(v) for v in parts[6:10])
            except ValueError:
                continue
            if conf < OCR_MIN_CONF or not _MEANINGFUL.search(parts[11]):
                continue
            cx, cy = (x0 + left + w / 2) * scale, ph - (y0 + top + h / 2) * scale
            key = (parts[11].strip(), round(cx), round(cy))
            if key in seen:
                continue
            seen.add(key)
            out.append(TextRecord(text=parts[11].strip(), x=round(cx, 2), y=round(cy, 2),
                                  height=round(h * scale, 2), layer="", kind="OCR",
                                  space=f"第 {page.number + 1} 页", owner=f"置信度 {conf:.0f}"))
    return out


def _pdf_tables(page: Any, words: list[TextRecord]) -> list[list[list[str]]]:
    """PDF 矢量线框表格：水平/竖直线段 → 网格 → 文字落格。"""
    from .gtable import HVLine
    ph = float(page.rect.height)
    hs, vs = [], []
    try:
        drawings = page.get_drawings()
    except Exception:
        return []
    for d in drawings:
        for item in d.get("items", []):
            if item[0] == "l":
                p1, p2 = item[1], item[2]
                x1, y1, x2, y2 = p1.x, ph - p1.y, p2.x, ph - p2.y
                if abs(y1 - y2) < 0.3 and abs(x1 - x2) > 2:
                    hs.append(HVLine(min(x1, x2), max(x1, x2), (y1 + y2) / 2, True))
                elif abs(x1 - x2) < 0.3 and abs(y1 - y2) > 2:
                    vs.append(HVLine(min(y1, y2), max(y1, y2), (x1 + x2) / 2, False))
            elif item[0] == "re":
                r = item[1]
                x0, y0, x1, y1 = r.x0, ph - r.y1, r.x1, ph - r.y0
                hs += [HVLine(x0, x1, y0, True), HVLine(x0, x1, y1, True)]
                vs += [HVLine(y0, y1, x0, False), HVLine(y0, y1, x1, False)]
    if len(hs) < 3 or len(vs) < 3 or not words:
        return []
    texts = [{"text": w.text, "x": w.x, "y": w.y, "h": max(w.height, 1.0),
              "w": _text_width(w.text, max(w.height, 1.0))} for w in words]
    min_h = max(min(t["h"] for t in texts), 1.0)
    out = []
    for comp_h, comp_v in _line_components(hs, vs, tol=1.0):
        if len(comp_h) < 3 or len(comp_v) < 3:
            continue
        x0 = min(min(h.start for h in comp_h), min(v.coord for v in comp_v))
        x1 = max(max(h.end for h in comp_h), max(v.coord for v in comp_v))
        y0 = min(min(h.coord for h in comp_h), min(v.start for v in comp_v))
        y1 = max(max(h.coord for h in comp_h), max(v.end for v in comp_v))
        sub = [t for t in texts if x0 <= t["x"] <= x1 and y0 <= t["y"] <= y1]
        if len(sub) < 4:
            continue
        grid = _grid_from(comp_h, comp_v, sub, min_h)
        if not grid:
            continue
        matrix = _compact([[c.text_content for c in row] for row in grid])
        if sum(1 for row in matrix for c in row if c) >= 4 and _fill_ratio(matrix) >= PDF_TABLE_MIN_FILL:
            out.append(matrix)
    return out


PDF_TABLE_MIN_FILL = 0.25   # 线框里有字的格子不到 1/4，多半是电路图而不是表格


def _line_components(hs: list, vs: list, tol: float) -> list[tuple[list, list]]:
    """横竖线按相交关系分组（并查集），一组即一张候选表格，避免整页线条拼成一张大网格。"""
    import bisect
    n_h = len(hs)
    parent = list(range(n_h + len(vs)))

    def find(i: int) -> int:
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    order = sorted(range(n_h), key=lambda i: hs[i].coord)
    coords = [hs[i].coord for i in order]
    for j, v in enumerate(vs):
        lo = bisect.bisect_left(coords, v.start - tol)
        hi = bisect.bisect_right(coords, v.end + tol)
        for k in range(lo, hi):
            h = hs[order[k]]
            if h.start - tol <= v.coord <= h.end + tol:
                a, b = find(order[k]), find(n_h + j)
                if a != b:
                    parent[a] = b
    groups: dict[int, tuple[list, list]] = {}
    for i, h in enumerate(hs):
        groups.setdefault(find(i), ([], []))[0].append(h)
    for j, v in enumerate(vs):
        groups.setdefault(find(n_h + j), ([], []))[1].append(v)
    return list(groups.values())


def scan_pdf(path: str, use_ocr: bool = True) -> NativeInventory:
    import pymupdf
    inv = NativeInventory(source=path, file_type="pdf", units="pt")
    doc = pymupdf.open(path)
    text_pages = ocr_pages = 0
    for page in doc:
        ph = float(page.rect.height)
        words = [TextRecord(text=w[4], x=round((w[0] + w[2]) / 2, 2), y=round(ph - (w[1] + w[3]) / 2, 2),
                            height=round(w[3] - w[1], 2), layer="", kind="PDF_TEXT",
                            space=f"第 {page.number + 1} 页")
                 for w in page.get_text("words") if str(w[4]).strip()]
        if words:
            text_pages += 1
        elif use_ocr:
            words = _ocr_page(page)
            if words:
                ocr_pages += 1
        inv.texts.extend(words)
        for matrix in _pdf_tables(page, words):
            inv.tables.append({"index": len(inv.tables) + 1, "source": "PDF 线框表格（候选，需核对）",
                               "block": "", "space": f"第 {page.number + 1} 页", "x": 0, "y": 0,
                               "rows": len(matrix), "cols": max(len(r) for r in matrix), "cells": matrix})
    inv.read_route.append(f"PDF 文字层：{text_pages}/{len(doc)} 页有文字")
    if ocr_pages:
        inv.read_route.append(f"OCR 兜底：{ocr_pages} 页（tesseract，需人工核对）")
    if not text_pages:
        inv.warnings.append("PDF 没有文字层（文字已转曲）"
                            + ("，已用 OCR 兜底，准确率有限" if ocr_pages else
                               "，本机未安装 tesseract，只能依赖视觉模型"))
    inv.coverage = [{"type": "PDF 页", "count": len(doc), "handling": "文字层 / 矢量线 / OCR"}]
    return inv


def scan_drawing(path: str, use_ocr: bool = True) -> NativeInventory:
    low = path.lower()
    if low.endswith((".dwg", ".dxf")):
        return scan_cad(path)
    if low.endswith(".pdf"):
        return scan_pdf(path, use_ocr=use_ocr)
    raise ValueError(f"不支持的图纸格式：{os.path.basename(path)}")


_CIRCUIT_REF = re.compile(r"^\s*([0-9A-Za-z][0-9A-Za-z\-/_.]{1,19})\s*[:：]\s*([A-Za-z]{0,3}\d{1,3}[A-Za-z]?)\s*$")
_BOX_WORDS = ("箱", "柜", "屏")
_SENTENCE = re.compile(r"[，。；,;！？]")
_RATING = re.compile(r"^\d+(\.\d+)?(A|AH|KA|KVA|KW|V|KV|MM2?)$")   # 40AH、630A、100KVA 是额定值不是箱号


def native_box_inventory(items: Iterable[tuple[str, str]]) -> dict[str, dict[str, Any]]:
    """从图纸原生文字里列出“有哪些配电箱/柜”，不依赖系统图版式规则。

    三类线索（都是图上写明的字，互相印证）：
      - caption：系统图标题“箱号 + 名称”，名称里带 箱/柜/屏，如“02AP01 配电总箱1”
      - symbol：平面图配电箱图例的属性值正好是箱号（天正设备块常见）
      - plan_circuits：平面图回路标注“箱号:回路号”，如“02AL01:N3”（只含去负载的回路，
        进线、备用、浪涌不在平面图上，所以是回路数下限）
    items 为 (文字, 对象类型)；返回 {箱号: {name, captions, symbols, plan_circuits, level}}，
    level：confirmed = 有系统图标题；plan = 只在平面图出现（需人工确认是否报价）。
    """
    from .cad_extractor import extract_panel_code

    inv: dict[str, dict[str, Any]] = {}

    def entry(code: str) -> dict[str, Any]:
        return inv.setdefault(code, {"code": code, "name": "", "captions": 0, "symbols": 0,
                                     "plan_circuits": set()})

    for text, kind in items:
        s = str(text or "").strip()
        if not s or len(s) > 60:
            continue
        m = _CIRCUIT_REF.match(s)
        if m:
            code = extract_panel_code(m.group(1))
            if code and code == m.group(1).upper():
                entry(code)["plan_circuits"].add(m.group(2).upper())
            continue
        code = extract_panel_code(s)
        if not code:
            continue
        if kind == "ATTRIB" and s.upper() == code:
            entry(code)["symbols"] += 1
            continue
        if _RATING.match(code):
            continue
        if s.upper().startswith(code):
            rest = s[len(code):].strip(" ：:-—")
            # 名称要像名称：不是说明句子、不以“#”序号开头（如“1#变压器柜”）、不过长
            if (rest and len(rest) <= 16 and any(w in rest for w in _BOX_WORDS)
                    and not _SENTENCE.search(rest) and not rest.startswith(("#", "＃"))):
                e = entry(code)
                e["captions"] += 1
                if not e["name"] or len(rest) < len(e["name"]):
                    e["name"] = rest
    for e in inv.values():
        e["level"] = "confirmed" if e["captions"] else "plan"
        e["plan_circuits"] = sorted(e["plan_circuits"], key=lambda c: (len(c), c))
    return {k: v for k, v in inv.items() if v["captions"] or v["symbols"] or v["plan_circuits"]}


def extraction_supplements(doc: Any) -> tuple[list[dict[str, Any]], list[Any]]:
    """给箱柜/回路识别（cad_extractor）用：现有读取漏掉的文字和表格网格。

    现有读取覆盖模型空间 TEXT / MTEXT 与块参照（块内文字 + 属性，单次插入）。这里补：
      - 表格 ACAD_TABLE：读表格图形块（*T…）里的文字（转换器常把表格对象本身读坏）
      - 多重引线文字
      - 代理对象（天正 / 浩辰等自定义对象）炸开代理图形后的文字（天正解密由 _Walker 完成）
      - MINSERT 阵列块除第一个之外的其余副本
    尺寸标注的显示文字不补：多为纯数字尺寸，混入会干扰回路编号与规格识别。

    返回 (文字列表, 表格网格列表)。文字为模型空间坐标的 {text, x, y, height, kind, owner, handle, layer}；
    表格为 gtable.StructuredTable，可直接交给 parse_switchgear_table。
    """
    from .gtable import StructuredTable

    walker = _Walker(doc)
    msp = doc.modelspace()
    texts: list[dict[str, Any]] = []
    tables: list[Any] = []
    for e in msp:
        kind = e.dxftype()
        try:
            if kind in ("ACAD_TABLE", "MULTILEADER", "MLEADER", "ACAD_PROXY_ENTITY"):
                texts.extend(walker._entity(e, 0))   # 代理对象文字在 _Walker 里已做天正解密
            elif kind == "INSERT":
                rows = int(e.dxf.get("row_count", 1) or 1)
                cols = int(e.dxf.get("column_count", 1) or 1)
                if rows * cols > 1:
                    local = walker.local(str(e.dxf.name))
                    for ins in list(e.multi_insert())[1:]:
                        texts.extend(walker._transformed(local, ins.matrix44(), str(e.dxf.name)))
        except Exception:
            continue

    for e in msp.query("ACAD_TABLE"):
        try:
            name = e.get_block_name()
            loc = e.get_insert_location()
            grid = _table_grid(doc, walker, name)
        except Exception:
            continue
        if not grid:
            continue
        ox, oy = float(loc[0]), float(loc[1])
        for row in grid:
            for c in row:
                c.min_x += ox
                c.max_x += ox
                c.min_y += oy
                c.max_y += oy
        xs = [c.min_x for r in grid for c in r] + [c.max_x for r in grid for c in r]
        ys = [c.min_y for r in grid for c in r] + [c.max_y for r in grid for c in r]
        tables.append(StructuredTable(num_rows=len(grid), num_cols=len(grid[0]), cells=grid,
                                      bbox=(min(xs), min(ys), max(xs), max(ys))))
    return texts, tables


def _table_grid(doc: Any, walker: _Walker, block_name: str):
    """表格图形块 → GTable 网格（局部坐标）。没有网格线时返回 None。"""
    from .gtable import extract_hv_lines
    blk = doc.blocks.get(block_name)
    if blk is None:
        return None
    texts = [{"text": it["text"], "x": it["x"], "y": it["y"], "h": float(it["height"] or 2.5),
              "w": _text_width(it["text"], float(it["height"] or 2.5))}
             for it in walker.local(block_name)]
    if not texts:
        return None
    min_h = max(min(t["h"] for t in texts), 0.5)
    lines = [e for e in blk if e.dxftype() in ("LINE", "LWPOLYLINE", "POLYLINE")]
    h, v = extract_hv_lines(lines)
    if len(h) < 2 or len(v) < 2:
        return None
    return _grid_from(h, v, texts, min_h)


def supplement_cad_texts(path: str, cad_texts: list[dict[str, Any]], inv: NativeInventory | None = None
                         ) -> list[dict[str, Any]]:
    """把现有 cad_texts 没读到的文字（表格、标注、多重引线、代理对象、块内文字）补进去。

    现有 extract_cad_entities 只读模型空间 TEXT / MTEXT / 块属性；按 (文字, 坐标取整) 去重后追加，
    不改动已有条目，供交叉核对与图签识别使用。
    """
    inv = inv or scan_cad(path)
    # 现有条目：模型空间顶层 TEXT / MTEXT 与顶层块的属性，用对象句柄去重（坐标基准不同，不能按坐标比）
    known_handles = {str((t.get("source") or {}).get("handle", "")) for t in cad_texts}
    seen: set = set()
    out = list(cad_texts)
    for t in inv.texts:
        if t.space != _SPACE_MODEL:
            continue
        if t.kind in ("TEXT", "MTEXT") and not t.owner:
            continue
        if t.kind == "ATTRIB" and t.handle in known_handles:
            continue
        key = (t.kind, t.text.strip(), round(t.x), round(t.y))
        if key in seen:
            continue
        seen.add(key)
        out.append({"type": t.kind, "text": t.text, "x": round(t.x, 2), "y": round(t.y, 2),
                    "height": t.height, "layer": t.layer, "source": {"owner": t.owner, "handle": t.handle,
                                                                   "origin": "native_takeoff"}})
    return out


# ---------------------------------------------------------------- 导出

def build_native_workbook(inv: NativeInventory, out_path: str) -> str:
    """原生读取结果导出 Excel：读取概览 / 文字清单 / 表格 / 图块统计 / 图层线长 / 字体检查。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F3A68")
    title_font = Font(bold=True, size=13)

    def sheet(wb, title, headers, rows, widths):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(vertical="center")
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        return ws

    wb = Workbook()
    ws = wb.active
    ws.title = "读取概览"
    s = inv.summary()
    ws.append(["图纸原生读取概览"])
    ws["A1"].font = title_font
    for k, v in [("文件", s["source"]), ("格式", s["file_type"].upper()),
                 ("读取路线", "；".join(s["read_route"])), ("图纸单位", inv.units),
                 ("文字条数", s["texts"]),
                 ("文字来源", "；".join(f"{k}×{v}" for k, v in s["texts_by_kind"].items())),
                 ("表格", s["tables"]), ("图块种类 / 引用总数", f"{s['block_kinds']} / {s['block_instances']}"),
                 ("有线长的图层", s["layers_with_length"])]:
        ws.append([k, v])
    ws.append([])
    ws.append(["提示"])
    ws[f"A{ws.max_row}"].font = Font(bold=True)
    for w in inv.warnings or ["无"]:
        ws.append(["", w])
    ws.append([])
    ws.append(["对象类型", "数量", "处理方式"])
    ws[f"A{ws.max_row}"].font = Font(bold=True)
    for c in inv.coverage:
        ws.append([c["type"], c["count"], c["handling"]])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 40

    sheet(wb, "文字清单", ["序号", "文字", "来源", "图层", "所在块 / 表格", "X", "Y", "字高", "空间"],
          [[i, t.text, t.kind, t.layer, t.owner, t.x, t.y, t.height, t.space]
           for i, t in enumerate(inv.texts, 1)],
          [7, 48, 11, 18, 24, 14, 14, 8, 14])

    wt = wb.create_sheet("表格")
    if not inv.tables:
        wt.append(["图纸中没有可还原的表格"])
    for t in inv.tables:
        wt.append([f"表格 {t['index']}（{t['source']}，{t['space']}，{t['rows']} 行 × {t['cols']} 列，"
                   f"位置 {t['x']}, {t['y']}）"])
        wt[f"A{wt.max_row}"].font = Font(bold=True)
        for row in t["cells"]:
            wt.append(row)
        wt.append([])
    for i in range(1, 16):
        wt.column_dimensions[get_column_letter(i)].width = 16

    sheet(wb, "图块统计", ["序号", "块名", "数量", "图层", "块内文字", "属性值", "块尺寸", "天正对象", "匿名块"],
          [[i, b["name"], b["count"], b["layers"], b["sample_text"], b["attributes"], b["size"],
            "是" if b["tangent"] else "", "是" if b["anonymous"] else ""]
           for i, b in enumerate(inv.blocks, 1)],
          [7, 30, 8, 22, 40, 36, 14, 9, 8])

    unit_m = 0.001 if inv.units in ("毫米", "未设置（按毫米折算）") else (1.0 if inv.units == "米" else None)
    sheet(wb, "图层线长", ["图层", "线段数", f"长度（{inv.units}）", "折合米", "构成"],
          [[r["layer"], r["count"], r["length"], round(r["length"] * unit_m, 2) if unit_m else "",
            r["types"]] for r in inv.layer_lengths],
          [26, 10, 18, 12, 40])

    sheet(wb, "字体检查", ["文字样式", "字体", "字体可用", "大字体", "大字体可用", "使用次数"],
          [[f["style"], f["font"], "是" if f["font_found"] else "缺失", f["bigfont"],
            ("是" if f["bigfont_found"] else "缺失") if f["bigfont"] else "", f["used"]]
           for f in inv.fonts],
          [24, 22, 10, 18, 10, 10])

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    return out_path


def inventory_cached(src_path: str, cache_json: str) -> dict[str, Any]:
    """按源文件修改时间缓存读取结果（JSON）。"""
    if os.path.isfile(cache_json) and os.path.getmtime(cache_json) >= os.path.getmtime(src_path):
        try:
            return json.load(open(cache_json, encoding="utf-8"))
        except Exception:
            pass
    data = scan_drawing(src_path).to_json()
    with open(cache_json, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False)
    return data


def inventory_from_json(data: dict[str, Any]) -> NativeInventory:
    inv = NativeInventory(**{k: v for k, v in data.items() if k not in ("summary", "texts")})
    inv.texts = [TextRecord(**t) for t in data.get("texts", [])]
    return inv
