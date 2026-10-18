# -*- coding: utf-8 -*-
"""成套设备组价两表模板：组价汇总 + 组价明细。

版面口径来自客户提供的成套设备报价单（后两张表），不是复制其中的数据：
本模块只接收结构化组价模型（箱体、器件行、四项费用），负责排版与金额落位。

口径要点（按客户要求定）：
- **没有封面、没有 1~6 行抬头**：不写报价日期、文件编号、报价单位、联系人电话；
  这是 AI 组价结果，不是供应商报价单。
- 标题只写表名，第一行就是表名，第二、三行是项目单位与项目名称。
- 待核对项不单独占页签，只在汇总表报价说明后追加一行提示。

模型结构::

    {
      "project": "项目名称", "owner": "项目单位",
      "currency": "人民币元", "category": "配电箱",
      "footer_note": "报价说明正文", "quote_by": "", "review_by": "",
      "boxes": [{
          "code": "02AL01", "name": "照明配电箱", "model": "", "location": "锅炉房",
          "quantity": 1,
          "items": [{"name": "", "spec": "", "unit": "只", "qty": 1,
                     "unit_price": 0.0, "brand": "", "note": ""}],
          "subtotal": 0.0, "auxiliary": 0.0, "manufacture": 0.0, "tax_total": 0.0,
          "unit_price": 0.0, "total_price": 0.0,
      }],
      "total_qty": 0, "total_price": 0.0, "grand_total": 0, "total_upper": "",
      "unresolved_count": 0,
    }
"""
from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import delivery as _delivery
from .quote_calc import rmb_upper

_CFG = _delivery().get("quote_template") or {}
_PALETTE = _delivery()["palette"]

FONT = _CFG.get("font", "宋体")
MONEY_FMT = "0.00_ "
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")

_thin = Side(style="thin", color=_PALETTE.get("border_dark", "595959"))
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

CATEGORY_FILL = PatternFill("solid", fgColor=_CFG.get("category_fill", "CCFFCC"))
CATEGORY_FILL_DETAIL = PatternFill("solid", fgColor=_CFG.get("category_fill_detail", "C6EFCE"))
BOX_FILL = PatternFill("solid", fgColor=_CFG.get("box_fill", "CCFFFF"))
HDR_FILL = PatternFill("solid", fgColor=_CFG.get("header_fill", "F2F2F2"))
TOTAL_FILL = PatternFill("solid", fgColor=_CFG.get("total_fill", "75923C"))

HEADERS = ["序号", "元件名称", "型号规格", "单位", "数量", "单价", "总价", "生产厂家", "备注"]
SUMMARY_HEADERS = ["序号", "柜号", "箱柜名称", "箱柜型号", "单位", "数量", "单价", "总价", "备注"]

SUMMARY_WIDTHS = _CFG.get("summary_widths") or [5.5, 15.6, 24.3, 15.6, 5.1, 5.1, 12.6, 14.0, 14.4]
DETAIL_WIDTHS = _CFG.get("detail_widths") or [5.0, 22.5, 26.2, 5.1, 5.1, 12.5, 13.5, 15.0, 13.1]

# 表头占 6 行：表名 / 项目单位 / 项目名称 / 金额单位 / 分类条 / 列头
HEADER_ROWS = 6
FIRST_DATA_ROW = HEADER_ROWS + 1

def _money(cell, value: Any = None):
    if value is not None:
        cell.value = value
    cell.number_format = MONEY_FMT
    cell.alignment = RIGHT
    return cell


def _header_block(ws, model: dict, title: str, category_fill) -> None:
    """表名 / 项目单位 / 项目名称 / 金额单位 / 分类条。

    没有报价日期、文件编号、报价单位、联系人电话——AI 组价结果不冒充供应商报价单。
    """
    ws.merge_cells("A1:I1")
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(name=FONT, size=16, bold=True, color=_PALETTE.get("dark_bold_text", "1F497D"))
    c.alignment = CENTER
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:I2")
    c = ws.cell(row=2, column=1, value=f"项目单位：{model.get('owner', '')}")
    c.font = Font(name=FONT, size=11)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws.merge_cells("A3:I3")
    c = ws.cell(row=3, column=1, value=f"项目名称：{model.get('project', '')}")
    c.font = Font(name=FONT, size=11)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws.merge_cells("H4:I4")
    c = ws.cell(row=4, column=8, value=f"金额单位：{model.get('currency', '人民币元')}")
    c.font = Font(name=FONT, size=11)
    c.alignment = RIGHT

    ws.merge_cells("A5:I5")
    c = ws.cell(row=5, column=1, value=model.get("category", "配电箱"))
    c.font = Font(name=FONT, size=11, bold=True)
    c.fill = category_fill
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in range(1, 10):
        ws.cell(row=5, column=col).border = BORDER
    ws.row_dimensions[5].height = 18


def _header_row(ws, row: int, headers: list[str]) -> None:
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(name=FONT, size=11)
        c.fill = HDR_FILL
        c.alignment = CENTER
        c.border = BORDER
        c.number_format = "@"
    ws.row_dimensions[row].height = 18


def _summary_sheet(ws, model: dict, detail_title: str) -> None:
    for j, w in enumerate(SUMMARY_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    _header_block(ws, model, _CFG.get("summary_title", "成套设备组价汇总"), CATEGORY_FILL)
    _header_row(ws, HEADER_ROWS, SUMMARY_HEADERS)

    start = FIRST_DATA_ROW
    r = start
    for i, box in enumerate(model.get("boxes", []), start=1):
        seq = ws.cell(row=r, column=1, value=f'=HYPERLINK("#{detail_title}!A{box["_row"]}:I{box["_end_row"]} ","1-{i}")')
        seq.font = Font(name=FONT, size=11, color=_CFG.get("link_color", "FF0000FF"))
        seq.alignment = CENTER
        seq.border = BORDER

        code_cell = ws.cell(row=r, column=2, value=box.get("code", ""))
        code_cell.number_format = "@"
        code_cell.alignment = LEFT

        vals = {
            3: box.get("name", ""),
            4: box.get("model", ""),
            5: box.get("unit", "台"),
            6: box.get("quantity", 1),
            7: box.get("unit_price", 0.0),
            8: box.get("total_price", 0.0),
            9: box.get("location", ""),
        }
        for col, v in vals.items():
            c = ws.cell(row=r, column=col, value=v)
            c.font = Font(name=FONT, size=11)
            c.alignment = LEFT if col in (3, 4, 9) else CENTER
            if col in (7, 8):
                _money(c)
        for col in range(1, 10):
            ws.cell(row=r, column=col).border = BORDER
        ws.row_dimensions[r].height = 18
        r += 1

    last = r - 1

    # 合计行（2 位小数）
    ws.cell(row=r, column=2, value="合计").font = Font(name=FONT, size=11, bold=True)
    ws.cell(row=r, column=2).alignment = LEFT
    ws.cell(row=r, column=5, value="台").font = Font(name=FONT, size=11)
    ws.cell(row=r, column=5).alignment = CENTER
    c = ws.cell(row=r, column=6, value=model.get("total_qty", 0))
    c.font = Font(name=FONT, size=11)
    c.alignment = CENTER
    c = ws.cell(row=r, column=8, value=model.get("total_price", 0.0))
    c.font = Font(name=FONT, size=11)
    _money(c)
    for col in range(1, 10):
        ws.cell(row=r, column=col).border = BORDER
    r += 1

    # 总计行（取整到元 + 大写金额）
    grand_total = model.get("grand_total")
    if grand_total is None:
        grand_total = int(round(float(model.get("total_price") or 0)))
    ws.cell(row=r, column=2, value="总计").font = Font(name=FONT, size=14, bold=True)
    c = ws.cell(row=r, column=3, value=int(grand_total))
    c.font = Font(name=FONT, size=14, bold=True)
    c.number_format = "0"
    c.alignment = CENTER
    ws.cell(row=r, column=5, value="台").font = Font(name=FONT, size=14, bold=True)
    ws.cell(row=r, column=5).alignment = CENTER
    ws.cell(row=r, column=6, value=model.get("total_qty", 0)).font = Font(name=FONT, size=14, bold=True)
    ws.cell(row=r, column=6).alignment = CENTER
    c = ws.cell(row=r, column=7, value=model.get("total_upper") or rmb_upper(grand_total))
    c.font = Font(name=FONT, size=11, bold=True)
    c.alignment = CENTER
    for col in range(1, 10):
        ws.cell(row=r, column=col).fill = TOTAL_FILL
        ws.cell(row=r, column=col).border = BORDER
    r += 1

    # 报价说明
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    c = ws.cell(row=r, column=1, value=_CFG.get("note_label", "报价说明："))
    c.font = Font(name=FONT, size=11)
    c.alignment = LEFT
    for col in range(1, 10):
        ws.cell(row=r, column=col).border = BORDER
    r += 1

    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
    c = ws.cell(row=r, column=2, value=model.get("footer_note", ""))
    c.font = Font(name=FONT, size=11)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[r].height = 30
    r += 1

    # 待核对项不单独出页签，只在报价说明后面说清有多少项没落定
    unresolved = int(model.get("unresolved_count") or 0)
    if unresolved:
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
        c = ws.cell(row=r, column=2,
                    value=(f"待核对：本表另有 {unresolved} 处图面存疑未确认，明细在系统内可查，"
                           "未确认前请勿直接对外报价。"))
        c.font = Font(name=FONT, size=11, bold=True, color=_PALETTE.get("warning_text", "C5221F"))
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 22
        r += 1

    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.cell(row=r, column=2, value=f"报价人：{model.get('quote_by', '')}").font = Font(name=FONT, size=11)
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
    ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=8)
    ws.cell(row=r, column=7, value=f"审核人：{model.get('review_by', '')}").font = Font(name=FONT, size=11)

    ws.print_area = f"A1:I{r}"


def _detail_sheet(ws, model: dict) -> None:
    for j, w in enumerate(DETAIL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    _header_block(ws, model, _CFG.get("detail_title", "成套设备组价明细"), CATEGORY_FILL_DETAIL)

    r = FIRST_DATA_ROW
    for i, box in enumerate(model.get("boxes", []), start=1):
        # 箱头：编号 / 柜号 / 型号 / 名称 / 备注
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
        head = {
            1: f"1-{i}",
            2: f"柜号：{box.get('code', '')}",
            3: f"型号：{box.get('model', '') or ''}",
            7: "名称：",
            8: box.get("name", ""),
            9: box.get("location", ""),
        }
        for col in range(1, 10):
            c = ws.cell(row=r, column=col)
            c.fill = BOX_FILL
            c.border = BORDER
            c.font = Font(name=FONT, size=11 if col == 2 else 10)
            c.alignment = CENTER if col in (1, 7) else LEFT
            if col in head:
                c.value = head[col]
        ws.row_dimensions[r].height = 18
        box["_row"] = r
        r += 1

        _header_row(ws, r, HEADERS)
        r += 1

        for seq, item in enumerate(box.get("items", []), start=1):
            ws.cell(row=r, column=1, value=seq).alignment = CENTER
            ws.cell(row=r, column=2, value=item.get("name", "")).alignment = LEFT
            ws.cell(row=r, column=3, value=item.get("spec", "")).alignment = LEFT
            ws.cell(row=r, column=4, value=item.get("unit", "只")).alignment = CENTER
            ws.cell(row=r, column=5, value=item.get("qty", 1)).alignment = CENTER
            c = ws.cell(row=r, column=6, value=item.get("unit_price", 0.0))
            _money(c)
            c = ws.cell(row=r, column=7, value=f"=E{r}*F{r}")
            _money(c)
            ws.cell(row=r, column=8, value=item.get("brand", "")).alignment = LEFT
            ws.cell(row=r, column=9, value=item.get("note", "")).alignment = LEFT
            for col in range(1, 10):
                cell = ws.cell(row=r, column=col)
                cell.font = Font(name=FONT, size=11)
                cell.border = BORDER
                if col == 2:
                    cell.number_format = "@"
            ws.row_dimensions[r].height = 18
            r += 1

        def cost_row(label: str, value) -> None:
            nonlocal r
            ws.cell(row=r, column=1, value=seq_no[0]).alignment = CENTER
            ws.cell(row=r, column=2, value=label).alignment = LEFT
            c = ws.cell(row=r, column=7, value=value)
            _money(c)
            ws.cell(row=r, column=8, value="").alignment = LEFT
            for col in range(1, 10):
                cell = ws.cell(row=r, column=col)
                cell.font = Font(name=FONT, size=11)
                cell.border = BORDER
            ws.row_dimensions[r].height = 18
            seq_no[0] += 1
            r += 1

        seq_no = [len(box.get("items", [])) + 1]
        cost_row("小计", box.get("subtotal", 0.0))
        cost_row("辅料", box.get("auxiliary", 0.0))
        cost_row("成套制作费", box.get("manufacture", 0.0))
        cost_row("税费", box.get("tax_total", 0.0))
        cost_row("单台合计", box.get("unit_price", 0.0))

        ws.cell(row=r, column=2, value="总计").alignment = LEFT
        ws.cell(row=r, column=4, value="台").alignment = CENTER
        ws.cell(row=r, column=5, value=box.get("quantity", 1)).alignment = CENTER
        c = ws.cell(row=r, column=7, value=box.get("total_price", 0.0))
        _money(c)
        for col in range(1, 10):
            cell = ws.cell(row=r, column=col)
            cell.font = Font(name=FONT, size=11)
            cell.border = BORDER
        ws.row_dimensions[r].height = 18
        box["_end_row"] = r
        r += 1

    ws.print_area = f"A1:I{r - 1}"


def fill(wb, model: dict) -> Any:
    """把组价模型渲染进工作簿：第 1 张汇总、第 2 张明细（先明细后汇总以取行号）。"""
    summary_name = _CFG.get("summary_sheet", "组价汇总")
    detail_name = _CFG.get("detail_sheet", "组价明细")

    ws_detail = wb.create_sheet(detail_name)
    _detail_sheet(ws_detail, model)

    ws_summary = wb.active
    ws_summary.title = summary_name
    _summary_sheet(ws_summary, model, detail_name)

    wb.move_sheet(ws_detail, offset=1)
    return wb
